# -*- coding: utf-8 -*-
"""JSON içeriğini Excel çalışma kitabına döker: bölüm başına bir sayfa, kılavuz metni yok. Gerektirir: openpyxl."""
import math
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.hyperlink import Hyperlink

LACIVERT, YESIL, GRI = "13294B", "47D7AC", "EEF1F5"
F = "Arial"
f_norm = Font(name=F, size=10, color="222222")
f_kalin = Font(name=F, size=10, bold=True, color=LACIVERT)
f_baslik = Font(name=F, size=13, bold=True, color="FFFFFF")
f_ust = Font(name=F, size=10, bold=True, color="FFFFFF")
f_bag = Font(name=F, size=10, underline="single", color="0B5CAD")
dolgu = lambda hex_: PatternFill("solid", fgColor=hex_)
ince = Side(style="thin", color="C9CED6")
kenar = Border(left=ince, right=ince, top=ince, bottom=ince)
sar = Alignment(wrap_text=True, vertical="top")
orta = Alignment(wrap_text=True, vertical="center", horizontal="center")


def sayfa_adi(no, kisa):
    return re.sub(r"[\\/*?:\[\]]", "-", f"{no} {kisa}")[:31]


def yaz(ws, hucre, deger, font=f_norm, fill=None, hiza=sar, cerceve=True):
    c = ws[hucre]
    c.value, c.font, c.alignment = deger, font, hiza
    if fill:
        c.fill = dolgu(fill)
    if cerceve:
        c.border = kenar
    return c


def bag(ws, hucre, sayfa, metin):
    c = ws[hucre]
    c.value, c.font = metin, f_bag
    c.hyperlink = Hyperlink(ref=hucre, location="'" + sayfa.replace("'", "''") + "'!A1", display=metin)
    return c


def yukseklik(metinler, genislikler):
    satir = max(sum(max(1, math.ceil(len(p) / max(5, w * 0.85))) for p in str(t or "").split("\n")) for t, w in zip(metinler, genislikler))
    return min(360, max(18, satir * 14 + 4))


def kur(ws, baslik, sutun_sayisi):
    ws.sheet_view.showGridLines = False
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(2, sutun_sayisi))
    yaz(ws, "A1", baslik, f_baslik, LACIVERT, Alignment(vertical="center", indent=1), False)
    ws.row_dimensions[1].height = 26
    ws.page_setup.orientation, ws.page_setup.paperSize, ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = "landscape", 9, 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.oddFooter.center.text = "Sayfa &P / &N"


def uret(belge, ctx, hedef, bicim):
    wb = Workbook()
    ozet = wb.active
    ozet.title = "00_Özet"
    a = (belge.get("bolumler", {}).get("1.1", {}).get("alanlar") or {})
    m = belge.get("meta") or {}
    kur(ozet, f"FS-TS · {a.get('gel_id') or ''} — {a.get('baslik') or ''}", 6)
    for i, w in enumerate([12, 46, 10, 10, 16, 60], 1):
        ozet.column_dimensions[L(i)].width = w
    r = 3
    if ctx.get("logo"):
        try:
            from openpyxl.drawing.image import Image
            img = Image(ctx["logo"])
            img.height, img.width = 30, int(30 * img.width / img.height)
            ozet.add_image(img, "A3")
            ozet.row_dimensions[3].height = 30
            r = 5
        except Exception:
            pass
    for et, d in [("Sürüm", m.get("surum")), ("Tarih", m.get("tarih")), ("Hazırlayan", m.get("hazirlayan")), ("Durum", m.get("durum")),
                  ("Müşteri / proje", m.get("musteri")), ("Geliştirme türleri", ", ".join(ctx["tur_ad"].get(t, t) for t in belge.get("turler") or [])),
                  ("Profil", belge.get("profil"))]:
        yaz(ozet, f"A{r}", et, f_kalin, GRI)
        ozet.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
        yaz(ozet, f"B{r}", d or "—")
        r += 1
    r += 1
    for i, b in enumerate(["No", "Başlık", "Ağırlık", "Sınıf", "Durum", "Not"], 1):
        yaz(ozet, f"{L(i)}{r}", b, f_ust, LACIVERT, orta)
    r += 1
    for g in ctx["gorunum"]:
        ozet.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
        yaz(ozet, f"A{r}", f"{g['no']} {g['ad']}", f_kalin, GRI)
        r += 1
        for v in g["bolumler"]:
            gecerli = v["durum"] == "gecerli"
            if gecerli:
                bag(ozet, f"A{r}", sayfa_adi(v["no"], v["kisa"]), v["no"]).border = kenar
            else:
                yaz(ozet, f"A{r}", v["no"])
            yaz(ozet, f"B{r}", v["baslik"])
            yaz(ozet, f"C{r}", v["agirlik"], hiza=orta)
            yaz(ozet, f"D{r}", v["sinif"], f_kalin if v["sinif"] == "Kritik" else f_norm, YESIL if v["sinif"] == "Kritik" else None, orta)
            yaz(ozet, f"E{r}", "Yazıldı" if gecerli else "Geçerli değil", hiza=orta)
            yaz(ozet, f"F{r}", f"{len(v['satirlar'])} satır" if gecerli else v["gerekce"])
            ozet.row_dimensions[r].height = yukseklik([v["baslik"], "" if gecerli else v["gerekce"]], [46, 60])
            r += 1
    ozet.freeze_panes = "A3"

    for g in ctx["gorunum"]:
        for v in g["bolumler"]:
            if v["durum"] != "gecerli":
                continue
            ws = wb.create_sheet(sayfa_adi(v["no"], v["kisa"]))
            ws.sheet_properties.tabColor = LACIVERT if int(g["no"]) % 2 else YESIL
            n = max(2, len(v["sutunlar"]))
            veri = v["satirlar"]
            gen = []
            for j in range(n):
                uzun = max([len(str(v["sutunlar"][j])) + 3 if j < len(v["sutunlar"]) else 0] + [len(str(s[j])) * 0.6 for s in veri if j < len(s)] + [14])
                gen.append(min(48, int(uzun)))
            if v["alanlar"]:
                gen[0] = max(gen[0], 28)
                gen[1] = max(gen[1], 48)
            for j, w in enumerate(gen, 1):
                ws.column_dimensions[L(j)].width = w
            kur(ws, f"{v['no']}  {v['baslik']}", n)
            bag(ws, "A2", "00_Özet", "← Özet")
            r = 4
            for et, d in v["alanlar"]:
                yaz(ws, f"A{r}", et, f_kalin, GRI)
                ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=n)
                yaz(ws, f"B{r}", d or "—")
                for c in range(3, n + 1):
                    ws.cell(r, c).border = kenar
                ws.row_dimensions[r].height = yukseklik([d], [sum(gen[1:])])
                r += 1
            if v["alanlar"]:
                r += 1
            if veri:
                for j, b in enumerate(v["sutunlar"], 1):
                    yaz(ws, f"{L(j)}{r}", b, f_ust, LACIVERT, orta)
                ws.row_dimensions[r].height = 30
                ilk = r
                r += 1
                for s in veri:
                    for j, d in enumerate(s, 1):
                        yaz(ws, f"{L(j)}{r}", d)
                    ws.row_dimensions[r].height = yukseklik(s, gen)
                    r += 1
                ws.auto_filter.ref = f"A{ilk}:{L(len(v['sutunlar']))}{r - 1}"
            if v["mermaid"]:
                r += 1
                yaz(ws, f"A{r}", "Diyagram kaynağı (Mermaid)", f_kalin, GRI)
                ws.merge_cells(start_row=r + 1, start_column=1, end_row=r + 1, end_column=n)
                yaz(ws, f"A{r + 1}", v["mermaid"], Font(name="Consolas", size=9))
                ws.row_dimensions[r + 1].height = min(300, 14 * (v["mermaid"].count("\n") + 1) + 6)
            ws.freeze_panes = "A3"

    ws = wb.create_sheet("Ek A Denetim")
    kur(ws, "Ek A: Mekanik denetim özeti", 4)
    for j, w in enumerate([10, 14, 44, 100], 1):
        ws.column_dimensions[L(j)].width = w
    r = 3
    for satir in ctx["ozet"]:
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
        yaz(ws, f"A{r}", satir, cerceve=False)
        ws.row_dimensions[r].height = yukseklik([satir], [160])
        r += 1
    r += 1
    for j, b in enumerate(["Önem", "Kural", "Konum", "Bulgu"], 1):
        yaz(ws, f"{L(j)}{r}", b, f_ust, LACIVERT, orta)
    r += 1
    for b in ctx["bulgular"]:
        for j, d in enumerate([b["sev"], b["kural"], b["yol"], b["mesaj"]], 1):
            yaz(ws, f"{L(j)}{r}", d)
        ws.row_dimensions[r].height = yukseklik([b["mesaj"]], [100])
        r += 1
    if not ctx["bulgular"]:
        yaz(ws, f"A{r}", "Bulgu yok", cerceve=False)
    wb.save(hedef)
