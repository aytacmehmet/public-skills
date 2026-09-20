# -*- coding: utf-8 -*-
"""JSON içeriğini Word belgesine döker: kapak, sürüm/onay, içindekiler, bölümler, denetim eki. Gerektirir: python-docx."""
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor, Twips

LACIVERT, YESIL, GRI = "13294B", "47D7AC", "EEF1F5"
FONT = "Arial"
DIKEY_GENISLIK, YATAY_GENISLIK = 16.0, 26.1  # cm, kenar boşlukları düşülmüş


def yazi(run, boyut=9, kalin=False, italik=False, renk="222222"):
    run.font.name, run.font.size, run.font.bold, run.font.italic = FONT, Pt(boyut), kalin, italik
    run.font.color.rgb = RGBColor.from_string(renk)
    rpr = run._r.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.append(rf)
    for k in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rf.set(qn(k), FONT)
    return run


def golge(cell, hex_):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_)
    cell._tc.get_or_add_tcPr().append(shd)


def hucre(cell, metin, genislik, boyut=8.5, kalin=False, renk="222222", dolgu=None, orta=False):
    cell.width = Cm(genislik)
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = p.paragraph_format.space_after = Pt(1.5)
    if orta:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    parcalar = str(metin if metin not in (None, "") else "—").split("\n")
    for i, parca in enumerate(parcalar):
        r = yazi(p.add_run(parca), boyut, kalin, renk=renk)
        if i < len(parcalar) - 1:
            r.add_break()
    if dolgu:
        golge(cell, dolgu)


def alan(paragraf, kod, yer_tutucu=""):
    f = OxmlElement("w:fldSimple")
    f.set(qn("w:instr"), kod)
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = yer_tutucu
    r.append(t)
    f.append(r)
    paragraf._p.append(f)


def tablo(doc, basliklar, satirlar, toplam_cm, agirliklar=None, ilk_sutun_etiket=False):
    n = len(basliklar) if basliklar else len(satirlar[0])
    if not agirliklar:
        agirliklar = []
        for j in range(n):
            veri = max([len(str(s[j])) for s in satirlar if j < len(s)] or [0])
            uzun = [len(str(basliklar[j])) * 0.9 + 2 if basliklar else 0, min(60, veri) * 0.55, veri + 5 if veri <= 12 else 0]
            agirliklar.append(max(10, max(uzun)))
    top = sum(agirliklar)
    gen = [toplam_cm * a / top for a in agirliklar]
    t = doc.add_table(rows=(1 if basliklar else 0) + len(satirlar), cols=n)
    t.style, t.autofit = "Table Grid", False
    for j, w in enumerate(gen):  # tblGrid genişlikleri; LibreOffice ve Word hücre genişliğinden önce buna bakar
        t.columns[j].width = Cm(w)
    r0 = 0
    if basliklar:
        tr = t.rows[0]._tr
        trpr = tr.get_or_add_trPr()
        trpr.append(OxmlElement("w:tblHeader"))
        for j, b in enumerate(basliklar):
            hucre(t.cell(0, j), b, gen[j], kalin=True, renk="FFFFFF", dolgu=LACIVERT, orta=True)
        r0 = 1
    for i, s in enumerate(satirlar):
        for j in range(n):
            etiket = ilk_sutun_etiket and j == 0
            hucre(t.cell(r0 + i, j), s[j] if j < len(s) else "", gen[j], kalin=etiket, renk=LACIVERT if etiket else "222222", dolgu=GRI if etiket else None)
    return t


def baslik(doc, metin, seviye):
    p = doc.add_heading(level=seviye)
    yazi(p.add_run(metin), {1: 16, 2: 12, 3: 10.5}[seviye], True, renk=LACIVERT)
    p.paragraph_format.keep_with_next = True
    return p


def paragraf(doc, metin, boyut=10, italik=False, renk="222222", sonra=4):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(sonra)
    yazi(p.add_run(metin), boyut, italik=italik, renk=renk)
    return p


def ust_alt_bilgi(section, logo, metin, genislik_cm):
    section.header.is_linked_to_previous = section.footer.is_linked_to_previous = False
    p = section.header.paragraphs[0]
    for twips in (4680, 9360):  # Header stilinin varsayılan sekmelerini temizle
        p.paragraph_format.tab_stops.add_tab_stop(Twips(twips), WD_TAB_ALIGNMENT.CLEAR)
    p.paragraph_format.tab_stops.add_tab_stop(Cm(genislik_cm), WD_TAB_ALIGNMENT.RIGHT)
    if logo:
        p.add_run().add_picture(logo, height=Cm(0.55))
    yazi(p.add_run("\t" + metin), 8, renk="555555")
    ppr = p._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    alt = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", "8"), ("w:space", "4"), ("w:color", YESIL)):
        alt.set(qn(k), v)
    bdr.append(alt)
    sekmeler = ppr.find(qn("w:tabs"))  # şema sırası: pBdr, tabs'tan önce gelir
    sekmeler.addprevious(bdr) if sekmeler is not None else ppr.append(bdr)
    f = section.footer.paragraphs[0]
    f.alignment = WD_ALIGN_PARAGRAPH.CENTER
    yazi(f.add_run("Sayfa "), 8, renk="555555")
    alan(f, "PAGE", "1")
    yazi(f.add_run(" / "), 8, renk="555555")
    alan(f, "NUMPAGES", "1")


def uret(belge, ctx, hedef, bicim):
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = FONT, Pt(10)
    ayar = doc.settings.element
    guncelle = OxmlElement("w:updateFields")
    guncelle.set(qn("w:val"), "true")
    sonraki = next((ayar.find(qn(k)) for k in ("w:hdrShapeDefaults", "w:footnotePr", "w:endnotePr", "w:compat") if ayar.find(qn(k)) is not None), None)
    sonraki.addprevious(guncelle) if sonraki is not None else ayar.append(guncelle)
    zoom = ayar.find(qn("w:zoom"))
    if zoom is not None and zoom.get(qn("w:percent")) is None:
        zoom.set(qn("w:percent"), "100")

    a = (belge.get("bolumler", {}).get("1.1", {}).get("alanlar") or {})
    m = belge.get("meta") or {}
    kimlik, ad = a.get("gel_id") or "FS-TS", a.get("baslik") or ""

    s1 = doc.sections[0]
    s1.page_width, s1.page_height = Cm(21), Cm(29.7)
    s1.left_margin = s1.right_margin = Cm(2.5)
    s1.top_margin, s1.bottom_margin = Cm(2.5), Cm(2)
    s1.different_first_page_header_footer = True
    ust_alt_bilgi(s1, ctx.get("logo"), f"FS-TS  ·  {kimlik}", DIKEY_GENISLIK)

    # kapak
    p = doc.add_paragraph()
    p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(90), Pt(30)
    if ctx.get("logo"):
        p.add_run().add_picture(ctx["logo"], height=Cm(1.6))
    yazi(doc.add_paragraph().add_run("Fonksiyonel ve Teknik Spesifikasyon"), 26, True, renk=LACIVERT)
    yazi(doc.add_paragraph().add_run("FS-TS  ·  SAP Cloud ERP geliştirmesi"), 14, renk=LACIVERT)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(24)
    yazi(p.add_run(f"{kimlik}  –  {ad}"), 12)
    turler = ", ".join(ctx["tur_ad"].get(t, t) for t in belge.get("turler") or [])
    tablo(doc, None, [["Geliştirme ID", kimlik], ["Geliştirme başlığı", ad], ["Süreç alanı / modül", a.get("modul")], ["Müşteri / proje", m.get("musteri")],
                      ["Geliştirme türleri", turler], ["Sürüm", m.get("surum")], ["Tarih", m.get("tarih")], ["Hazırlayan", m.get("hazirlayan")],
                      ["Durum", m.get("durum")]], DIKEY_GENISLIK, [3, 7], True)
    doc.add_page_break()
    baslik(doc, "Sürüm geçmişi", 3)
    tablo(doc, ["Sürüm", "Tarih", "Yazan", "Değişiklik"], belge.get("surum_gecmisi") or [["—"] * 4], DIKEY_GENISLIK, [1.2, 2, 3, 6])
    doc.add_paragraph()
    baslik(doc, "Onaylar", 3)
    tablo(doc, ["Rol", "Ad soyad", "Tarih"], belge.get("onaylar") or [["—"] * 3], DIKEY_GENISLIK, [4, 4, 2])
    doc.add_page_break()
    baslik(doc, "İçindekiler", 3)
    alan(doc.add_paragraph(), 'TOC \\o "1-2" \\h \\z \\u', "İçindekiler, belge Word'de açıldığında alanlar güncellenince oluşur.")

    # gövde (yatay)
    s2 = doc.add_section(WD_SECTION.NEW_PAGE)
    s2.orientation = WD_ORIENT.LANDSCAPE
    s2.page_width, s2.page_height = Cm(29.7), Cm(21)
    s2.left_margin = s2.right_margin = Cm(1.8)
    s2.top_margin, s2.bottom_margin = Cm(2), Cm(1.8)
    s2.different_first_page_header_footer = False
    ust_alt_bilgi(s2, ctx.get("logo"), f"FS-TS  ·  {kimlik}", YATAY_GENISLIK)

    ilk = True
    for g in ctx["gorunum"]:
        h1 = baslik(doc, f"{g['no']} {g['ad']}", 1)
        h1.paragraph_format.page_break_before = not ilk
        ilk = False
        for v in g["bolumler"]:
            baslik(doc, f"{v['no']} {v['baslik']}", 2)
            if v["durum"] != "gecerli":
                paragraf(doc, f"Geçerli değil – {v['gerekce']}", italik=True, renk="666666")
                continue
            if v["alanlar"]:
                tablo(doc, None, [[e, d] for e, d in v["alanlar"]], YATAY_GENISLIK, [2.4, 10], True)
                doc.add_paragraph().paragraph_format.space_after = Pt(2)
            if v["satirlar"]:
                tablo(doc, v["sutunlar"], v["satirlar"], YATAY_GENISLIK)
                doc.add_paragraph().paragraph_format.space_after = Pt(2)
            if v["mermaid"]:
                paragraf(doc, "Diyagram kaynağı (Mermaid):", 9, renk="555555", sonra=1)
                for satir in v["mermaid"].split("\n"):
                    q = doc.add_paragraph()
                    q.paragraph_format.space_after = Pt(0)
                    r = q.add_run(satir)
                    r.font.name, r.font.size = "Consolas", Pt(8.5)

    h1 = baslik(doc, "Ek A: Mekanik denetim özeti", 1)
    h1.paragraph_format.page_break_before = True
    for satir in ctx["ozet"]:
        paragraf(doc, satir, 9.5, sonra=2)
    if ctx["bulgular"]:
        doc.add_paragraph()
        tablo(doc, ["Önem", "Kural", "Konum", "Bulgu"], [[b["sev"], b["kural"], b["yol"], b["mesaj"]] for b in ctx["bulgular"][:60]],
              YATAY_GENISLIK, [1.2, 1.8, 5, 12])
    paragraf(doc, "Gösterge yalnız mekanik kurallardan hesaplanır ve varsayılan parametrelere dayanır; kalite (LLM) puanını gerçek puanlama motoru ölçer.",
             8.5, italik=True, renk="666666")
    doc.core_properties.title = f"FS-TS {kimlik}"
    doc.save(hedef)
