#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Skill öz testi: türev dosyalar güncel mi, örnek içerik temiz mi, denetim hataları yakalıyor mu, dökümler üretiliyor mu.
Kullanım: python scripts/oz_test.py        (şablon, tanım ya da betik her değiştiğinde çalıştırın)"""
import copy
import json
import os
import subprocess
import sys
import tempfile
import zipfile

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import bv  # noqa: E402

KOK, sonuc = bv.KOK, []
ORNEK, PROJE = os.path.join(KOK, "assets", "ornek-icerik.json"), os.path.join(KOK, "assets", "proje-ornek.json")


def kontrol(ad, ok, ayrinti=""):
    sonuc.append(ok)
    print(("GEÇTİ " if ok else "KALDI ") + ad + (f"  [{ayrinti}]" if ayrinti else ""))


def calistir(*arg):
    return subprocess.run([sys.executable, os.path.join(KOK, "scripts", arg[0])] + list(arg[1:]), capture_output=True, text=True, encoding="utf-8", errors="replace")


r = calistir("derle.py", "--kontrol")
kontrol("Türev dosyalar tanım dosyasıyla güncel", r.returncode == 0, r.stdout.strip())
b = [x for x in bv.T["bolumler"]]
kontrol("Tanım: 32 bölüm, taban 119, bonus 35", len(b) == 32 and sum(x["agirlik"] for x in b if x["sinif"] != "Bonus") == 119
        and sum(x["agirlik"] for x in b if x["sinif"] == "Bonus") == 35)
kontrol("Tanım: 23 kural, 20 eski kural, 6 hazırlık kalemi", (len(bv.T["kurallar"]), len(bv.T["eski_kurallar"]), len(bv.T["hazirlik_kalemleri"])) == (23, 20, 6))

belge = bv.yukle(ORNEK)
bulgular, g, sayilan = bv.tam_denetim(belge, PROJE)
kontrol("Örnek içerik: bulgu yok", not bulgular, f"{len(bulgular)} bulgu, gösterge {g['sonuc']}")
bulgular2, _, _ = bv.tam_denetim(belge)
kontrol("Proje kataloğu olmadan standart nesneler UYD_001 verir", any(x["kural"] == "UYD_001" for x in bulgular2))
m62 = bv.turet_62(belge)
kontrol("6.2 türetimi: her REQ için tam zincir", len(m62) == 4 and all(s[7] == "Tam" for s in m62), str([s[7] for s in m62]))

bozuk = copy.deepcopy(belge)
bozuk["bolumler"]["3.5"]["satirlar"][1][3] = "Gerekli kontroller yapılır ve ilgili tablolardan okunur"
bozuk["bolumler"]["3.5"]["satirlar"][3][6] = "Zorunlu alan boş → MSG-44"
bozuk["bolumler"]["3.5"]["satirlar"][2][4] = "log tablosu"
bozuk["bolumler"]["3.4"]["satirlar"][0][8] = ""
bozuk["bolumler"]["2.2"]["alanlar"]["cozum_ozeti"] = bv.T["yer_tutucu"]
bozuk["bolumler"]["6.1"]["satirlar"] = [s for s in bozuk["bolumler"]["6.1"]["satirlar"] if s[1] == "Mutlu yol"]
bozuk["bolumler"]["3.3"] = {"gecerli": False, "gerekce": "deneme"}
bozuk["bolumler"]["4.1"]["satirlar"] = bozuk["bolumler"]["4.1"]["satirlar"][:-2]
kurallar = {x["kural"] for x in bv.tam_denetim(bozuk, PROJE)[0]}
beklenen = {"ALGO_003", "ZINCIR_001", "ALGO_001", "GEN_002", "MAP_001", "GEN_001", "TEST_002", "KRITIK_001", "SAP_001", "UYD_001", "ZINCIR_002"}
kontrol("Bozulmuş kopyada beklenen tüm kurallar yakalanıyor", beklenen <= kurallar, "eksik: " + ", ".join(sorted(beklenen - kurallar)))

acik_b = copy.deepcopy(belge)
s73 = [x for x in acik_b["bolumler"]["7.3"]["satirlar"] if x[5] == "Açık"]
s73[0][3], s73[1][4], s73[2][2] = "—", "—", "9.9"
b73 = [x for x in bv.tam_denetim(acik_b, PROJE)[0] if x["kural"] in ("OPEN_001", "OPEN_002")]
kontrol("7.3: sahipsiz satır kusur, tarihsiz satır girdi bekleyen, geçersiz bölüm numarası kusur",
        sorted((x["kural"], x["acik"]) for x in b73) == [("OPEN_001", False), ("OPEN_001", True), ("OPEN_002", False)], str([(x["kural"], x["acik"]) for x in b73]))

with tempfile.TemporaryDirectory() as d:
    r = calistir("bv.py", "iskelet", "--tur", "R", "--profil", "hafif", "--cikti", os.path.join(d, "i.json"))
    isk = bv.yukle(os.path.join(d, "i.json"))
    kontrol("İskelet: meta ve sürüm geçmişi YAPI_010 vermiyor", not [x for x in bv.tam_denetim(isk, None)[0] if x["kural"] == "YAPI_010"])
    kontrol("İskelet: Rapor/hafif için 3.2 ve 3.6 var; 3.3, 3.4, 4.4 ve bonuslar yok",
            {"3.2", "3.6"} <= set(isk["bolumler"]) and not {"3.3", "3.4", "4.4", "3.7", "3.9", "6.2"} & set(isk["bolumler"]), r.stdout.strip()[:80])
    yama = json.dumps([{"islem": "ayarla", "yol": "/bolumler/1.1/alanlar/gel_id", "deger": "GEL-TEST-01"},
                       {"islem": "ekle", "yol": "/bolumler/3.2/satirlar", "deger": ["SEL-01"] + ["—"] * 8}])
    calistir("bv.py", "yama", os.path.join(d, "i.json"), "--op", yama)
    isk = bv.yukle(os.path.join(d, "i.json"))
    kontrol("Yama: ayarla ve ekle çalışıyor", isk["bolumler"]["1.1"]["alanlar"]["gel_id"] == "GEL-TEST-01" and len(isk["bolumler"]["3.2"]["satirlar"]) == 1)

    basliklar = [f"{x['no']} {x['baslik']}" for x in bv.T["bolumler"] if bv.bolum_durumu(x, belge["turler"], belge["profil"]) != "profil_disi"]
    r = calistir("bv.py", "dok", ORNEK, "--proje", PROJE, "--bicim", "md,vault", "--cikti", d)
    md = next((os.path.join(d, f) for f in os.listdir(d) if f.endswith(".md")), None)
    metin = open(md, encoding="utf-8").read() if md else ""
    kontrol("Markdown dökümü: profilde açık tüm başlıklar var", bool(md) and all(f"### {t}" in metin for t in basliklar), r.stderr.strip()[:120])
    vault = next((os.path.join(d, f) for f in os.listdir(d) if f.endswith("_vault")), None)
    notlar = {x for _, _, fs in os.walk(vault or d) for f in fs for x in (f, os.path.splitext(f)[0])}
    import re
    kirik = [h for dp, _, fs in os.walk(vault or d) for f in fs if f.endswith(".md")
             for h in re.findall(r"\[\[([^\]|\\]+)", open(os.path.join(dp, f), encoding="utf-8").read()) if h not in notlar]
    kontrol("Vault dökümü: tüm [[bağlantılar]] çözülüyor", bool(vault) and not kirik, str(kirik[:3]))
    for bicim, modul in (("docx", "docx"), ("xlsx", "openpyxl")):
        try:
            __import__(modul)
        except ImportError:
            print(f"ATLANDI {bicim} dökümü ({modul} kurulu değil)")
            continue
        r = calistir("bv.py", "dok", ORNEK, "--proje", PROJE, "--bicim", bicim, "--cikti", d)
        yol = next((os.path.join(d, f) for f in os.listdir(d) if f.endswith("." + bicim)), None)
        ok = bool(yol) and zipfile.ZipFile(yol).testzip() is None
        if ok and bicim == "docx":
            import docx
            basl = [p.text for p in docx.Document(yol).paragraphs if p.style is not None and p.style.name.startswith("Heading")]
            ok = all(t in basl for t in basliklar)
        if ok and bicim == "xlsx":
            import openpyxl
            wb = openpyxl.load_workbook(yol)
            ok = "00_Özet" in wb.sheetnames and "3.5 İşlem mantığı" in wb.sheetnames and "3.2 Seçim ekranı" not in wb.sheetnames
        kontrol(f"{bicim} dökümü üretiliyor ve yapısı doğru", ok, r.stderr.strip()[:120])

print(f"\nSONUÇ: {sum(sonuc)}/{len(sonuc)}")
sys.exit(0 if all(sonuc) else 1)
