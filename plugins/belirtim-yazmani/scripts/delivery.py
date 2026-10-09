# SPDX-License-Identifier: GPL-3.0-only
"""A-F release gates. Private controls never enter the developer projection."""
import copy, io, json, re, tempfile, zipfile, uuid
from pathlib import Path
import bv2 as b
import source_pool
from jsonschema import Draft202012Validator

SCHEMA=json.loads((b.ROOT/'schema/handoff.schema.json').read_text(encoding='utf-8'))
SCHEMA_V2=json.loads((b.ROOT/'schema/handoff-v2.schema.json').read_text(encoding='utf-8'))
GAP=re.compile(r'(?i)KARAR BEKLİYOR|BİLGİ BEKLİYOR|DOĞRULANACAK|\bTBD\b|NEEDS[_ ]CLARIFICATION|\bUNKNOWN\b|\bsafe_default\b|uygun şekilde|\bgerekirse\b|\bvb\.|\bmümkünse\b|…')
CODE=re.compile(r'(?im)\bCAST\s*\(|\bCASE\s+WHEN\b|\bSELECT\s+.+\bFROM\b|\bdefine\s+(?:root\s+)?view\b|^\s*(?:REPORT\s+[zy]\w+\.|CLASS\s+.+(?:DEFINITION|IMPLEMENTATION)|METHOD\s+\w+\.|ENDMETHOD\.|DATA\s*\(|IF\s+.+\.|ENDIF\.)|```(?:abap|cds|bdef|js|javascript|typescript|python|pseudo(?:code)?)\b')
TOOL=re.compile(b.PRODUCER.pattern+r'|(?i:belirtim[- ]yazman[ıi]|spec[- ]writer|\bbv2?\.py\b|legacy_core\.py|@toon-format/toon|tool_receipt|tool_capabilities|execution_mode|mutation_policy|reasoning_effort|token_budget|obsidian|\[\[[^\]]+\]\]|\b(?:skill|plugin|mcp)[/-]|(?:^|\s)/plan\b)')
PRIVATE={'producerApproval','handoffApproval','approvalReceipt','onaylar','toolReceipts','tool_receipts','executionMode','mutationPolicy','llmModel','llm_model','developerModel','modelProvider','producerModel','providerSettings','reasoningEffort','reasoningLevel','tokenBudget','buildPrompt','systemPrompt','developerPrompt','agentName','toolCapabilities','producerMetadata'}
PRIVATE_NORMALIZED={re.sub('[^a-z0-9]','',k.lower()) for k in PRIVATE}
PRIVATE_NORMALIZED.update({'workprofile','workprofilerequired','episodeid','workassessment','workcapabilities','workcapacity','dispatchattempts',
                           'worktask','bywtask','coordinatorfile','modelsuggestion','preparationattempts','reviewattempts',
                           'nativedispatch','tooluseid','reviewrole','blindnesssha256'})
AI_INSTRUCTION=re.compile(r'(?im)^\s*(?:you are (?:a|an) [^\n]*(?:developer|assistant|agent)|(?:system|developer|build) prompt\s*:|(?:use|invoke|call) \$[a-z-]+)|ignore (?:previous|all) [^\n]*instructions')
COLLECTIONS=['requirements','acceptance_criteria','test_cases','objects','business_rules','screens','media','ui_callouts','ui_elements','ui_actions','messages','statuses','status_transitions','cds_fields','table_fields']
NETWORK=re.compile(r'(?i)https?://|fetch\s*\(|XMLHttpRequest|WebSocket\s*\(|EventSource\s*\(|sendBeacon\s*\(|importScripts\s*\(|import\s*\(|\beval\s*\(|\bFunction\s*\(|(?:src|href)\s*=\s*[\"\x27]\s*//|(?:url\s*\(|@import)[^;\n]*//')

def issue(code,path,message): return {'code':code,'pointer':path,'message':message}

def resolve(value,pointer):
    if pointer=='': return value
    if not isinstance(pointer,str) or not pointer.startswith('/'): raise b.Invalid('Use a decoded-data JSON Pointer')
    current=value
    for part in pointer[1:].split('/'):
        key=part.replace('~1','/').replace('~0','~')
        if isinstance(current,list):
            if not re.fullmatch('0|[1-9][0-9]*',key): raise b.Invalid('Invalid array index')
            current=current[int(key)]
        else: current=current[key]
    return current

def privacy(value,path=''):
    errors=[]
    if isinstance(value,dict):
        for k,v in value.items():
            if re.sub('[^a-z0-9]','',k.lower()) in PRIVATE_NORMALIZED: errors.append(issue('C5_PRIVATE',path+'/'+k,'Private control/producer field'))
            errors+=privacy(v,path+'/'+k)
    elif isinstance(value,list):
        for i,v in enumerate(value): errors+=privacy(v,path+'/'+str(i))
    elif isinstance(value,str):
        if TOOL.search(value): errors.append(issue('C7_ATTRIBUTION',path,'Producer/tool/model syntax'))
        if AI_INSTRUCTION.search(value): errors.append(issue('B7_AI_PROMPT',path,'Build/AI instruction is not developer specification content'))
        if CODE.search(value) or b.IMPLEMENTATION.search(value): errors.append(issue('B2_B4_CODE',path,'Implementation, expression or pseudo-code'))
    return errors

def schema_errors(spec):
    schema=SCHEMA if spec.get('schemaVersion')=='3.0' else SCHEMA_V2
    return [issue('D3_SCHEMA','/'+ '/'.join(map(str,e.path)),e.message) for e in Draft202012Validator(schema).iter_errors(spec)]

def _shape_and_profile_v2(spec):
    errors=schema_errors(spec)
    if errors: return errors
    errors+=privacy(spec)
    for path,value in b.flatten(spec).items():
        if isinstance(value,str) and GAP.search(value): errors.append(issue('D2_D4_GAP',path,'Unresolved or evasive content'))
    ids={}; labels={}; sets={}
    for name in COLLECTIONS:
        sets[name]={r['id'] for r in spec[name]}
        for i,row in enumerate(spec[name]):
            if row['id'] in ids: errors.append(issue('D7_DUPLICATE','/'+name+'/'+str(i),row['id']))
            ids[row['id']]=(name,row); labels[row['id']]='/'+name+'/'+str(i)
    def ref(target,kind,path):
        if target not in sets[kind]: errors.append(issue('D7_REF',path,str(target)))
    for i,r in enumerate(spec['requirements']):
        if not r['acceptance_criteria']: errors.append(issue('D7_AC','/requirements/'+str(i),'Acceptance criterion required'))
        for x in r['acceptance_criteria']:
            ref(x,'acceptance_criteria',labels[r['id']])
            if x in ids and ids[x][1].get('requirement_ref')!=r['id']: errors.append(issue('D7_BIDIRECTIONAL',labels[r['id']],x))
    for r in spec['acceptance_criteria']:
        ref(r['requirement_ref'],'requirements',labels[r['id']])
        if not r['test_cases']: errors.append(issue('D7_TEST',labels[r['id']],'Test required'))
        for x in r['test_cases']:
            ref(x,'test_cases',labels[r['id']])
            if x in ids and ids[x][1].get('acceptance_ref')!=r['id']: errors.append(issue('D7_BIDIRECTIONAL',labels[r['id']],x))
    for r in spec['test_cases']:
        ref(r['acceptance_ref'],'acceptance_criteria',labels[r['id']])
        if not r['object_refs']: errors.append(issue('D7_OBJECT',labels[r['id']],'Test must link implementation objects'))
        for x in r['object_refs']: ref(x,'objects',labels[r['id']])
    for r in spec['objects']:
        if not r['requirement_refs']: errors.append(issue('D7_ORPHAN_OBJECT',labels[r['id']],'Orphan object'))
        if r['operation'] in ('NEW','CHANGE') and r['owner_development']!=spec['meta']['development_id']: errors.append(issue('E1_FOREIGN_MUTATION',labels[r['id']],'Mutating another development requires its own handoff'))
        for x in r['requirement_refs']: ref(x,'requirements',labels[r['id']])
        try:
            if not re.fullmatch(spec['naming_rules']['object_pattern'],r['name']): errors.append(issue('A4_OBJECT_NAME',labels[r['id']],r['name']))
        except re.error: errors.append(issue('A4_PATTERN','/naming_rules/object_pattern','Invalid pattern'))
    for k,v,pat in [('name',spec['meta']['development_name'],'development_pattern'),('version',spec['meta']['handoff_version'],'version_pattern')]:
        try:
            if not re.fullmatch(spec['naming_rules'][pat],v): errors.append(issue('A4_NAMING','/meta/'+k,'Project naming violation'))
            b.safe_name(v)
        except (re.error,b.Invalid): errors.append(issue('A4_NAMING','/meta/'+k,'Unsafe name/version/pattern'))
    for r in spec['business_rules']:
        if not any(x['boundary'] for x in r['examples']): errors.append(issue('D3_BOUNDARY',labels[r['id']],'Boundary example required'))
        if r['multi_condition'] and not r['decision_table']: errors.append(issue('D3_DECISION_TABLE',labels[r['id']],'Decision table required'))
    for r in spec['messages']: ref(r['trigger_rule_ref'],'business_rules',labels[r['id']])
    for r in spec['status_transitions']:
        ref(r['from_ref'],'statuses',labels[r['id']]);ref(r['to_ref'],'statuses',labels[r['id']])
    for r in spec['statuses']:
        if not r['final'] and not any(t['from_ref']==r['id'] for t in spec['status_transitions']): errors.append(issue('D3_STATE_EXIT',labels[r['id']],'Non-final state has no exit'))
    for name in ['cds_fields','table_fields']:
        for r in spec[name]:
            ref(r['object_ref'],'objects',labels[r['id']])
            field='expression' if name=='cds_fields' else 'derivation'
            if r[field] is not None and re.search(r'[=+*/<>|&]|\b(?:IF|THEN|ELSE|CAST|SELECT|CASE)\b',r[field],re.I): errors.append(issue('B4_EXPRESSION',labels[r['id']]+'/'+field,'Use plain business text and input/output examples'))
    ui='U' in spec['meta']['types']
    if ui and any(not spec[n] for n in ['screens','media','ui_callouts','ui_elements','ui_actions']): errors.append(issue('B5_D3_UI','/screens','UI requires PNG/callout/element/action profile'))
    callout_keys=set()
    for r in spec['screens']:
        for x in r['media_refs']: ref(x,'media',labels[r['id']])
        for x in r['ui_element_ids']: ref(x,'ui_elements',labels[r['id']])
    for r in spec['media']:
        ref(r['screen_ref'],'screens',labels[r['id']])
        try: b.safe_name(r['path'])
        except b.Invalid: errors.append(issue('B5_PATH',labels[r['id']],'Unsafe image path'))
        if not r['path'].startswith('mockup/screens/') or not re.fullmatch('[a-f0-9]{64}',r['sha256']): errors.append(issue('B5_MEDIA',labels[r['id']],'PNG path/hash required'))
    for r in spec['ui_callouts']:
        ref(r['media_ref'],'media',labels[r['id']]);ref(r['ui_element_ref'],'ui_elements',labels[r['id']])
        key=(r['media_ref'],r['number'])
        if key in callout_keys: errors.append(issue('B5_CALLOUT_NUMBER',labels[r['id']],'Duplicate callout number'))
        callout_keys.add(key)
        if r['ui_element_ref'] in ids and ids[r['ui_element_ref']][1].get('label')!=r['label']: errors.append(issue('F4_LABEL',labels[r['id']],'Callout label differs'))
    defaults={r['id'] for r in spec['functional_defaults']['rules']}
    for r in spec['ui_elements']:
        ref(r['screen_ref'],'screens',labels[r['id']])
        if not any(c['ui_element_ref']==r['id'] for c in spec['ui_callouts']): errors.append(issue('B5_CALLOUT_COVERAGE',labels[r['id']],'Every visible element needs a callout'))
        if r['source']['kind']=='CDS_FIELD': ref(r['source']['ref'],'cds_fields',labels[r['id']])
        elif r['source']['ref'] is not None: errors.append(issue('D3_UI_ONLY',labels[r['id']],'UI-only has no CDS ref'))
        if r['default']['kind']=='PROJECT_DEFAULT' and r['default']['ref'] not in defaults: errors.append(issue('D5_DEFAULT_REF',labels[r['id']],'Missing project default'))
        if r['value_help']['kind']=='CDS_FIELD': ref(r['value_help']['ref'],'cds_fields',labels[r['id']])
        if r['value_help']['kind']=='FIXED' and not r['value_help']['values']: errors.append(issue('D3_VALUE_HELP',labels[r['id']],'Fixed value help empty'))
    kinds={'happy':'HAPPY','validation_error':'VALIDATION_ERROR','authorization_error':'AUTHORIZATION_ERROR'}
    for r in spec['ui_actions']:
        ref(r['screen_ref'],'screens',labels[r['id']])
        for path,kind in kinds.items():
            if not r['trace_paths'][path]: errors.append(issue('D3_TRACE_PATH',labels[r['id']],path))
            for tid in r['trace_paths'][path]:
                ref(tid,'test_cases',labels[r['id']])
                if tid in ids and ids[tid][1].get('kind')!=kind: errors.append(issue('D3_TRACE_KIND',labels[r['id']],path))
    for dep in spec['dependencies']:
        if dep['development_id']==spec['meta']['development_id']: errors.append(issue('E1_DEPENDENCY','/dependencies','Self dependency'))
    if not spec['scope_items']['included']: errors.append(issue('D3_SCOPE','/scope_items','Scope must be explicit'))
    if spec['meta']['mode']=='UPDATE':
        base=spec['baseline'];kind=base['kind']
        if kind=='SYSTEM' and not all(base[k] for k in ['system','client','verified_at','reference_sha256']) or kind=='FINAL_HANDOFF' and not all(base[k] for k in ['reference_name','reference_version','reference_sha256']): errors.append(issue('E4_BASELINE','/baseline','Incomplete baseline identity'))
        if kind in ('NEW','UNKNOWN') or not base['evidence_paths']: errors.append(issue('E3_BASELINE_UNKNOWN','/baseline','Verified presence/absence evidence required'))
    elif spec['baseline']['kind']!='NEW': errors.append(issue('E4_NEW_BASELINE','/baseline','New design needs NEW baseline'))
    if spec['changelog'][-1]['version']!=spec['meta']['handoff_version']: errors.append(issue('A5_CHANGELOG','/changelog','Current version missing'))
    return errors

def shape_and_profile(spec):
    if spec.get('schemaVersion')!='3.0': return _shape_and_profile_v2(spec)
    import handoff3 as h
    errors=schema_errors(spec)
    if errors: return errors
    return _shape_and_profile_v2(h.core_spec(spec))+privacy({k:spec[k] for k in h.EXTRAS})+h.profile(spec)

def get_state(doc):
    data=doc.get('delivery')
    if not isinstance(data,dict) or set(data)!={'spec','control'}: raise b.Invalid('AP2.v4 needs delivery.spec and private delivery.control; run release-init')
    if not isinstance(data['spec'],dict) or not isinstance(data['control'],dict): raise b.Invalid('Invalid delivery workspace')
    return data['spec'],data['control']

def revision(spec): return b.digest(b.canonical(spec))

def context_packet(spec,collections):
    groups={name:spec[name] for name in COLLECTIONS}
    if spec.get('schemaVersion')=='3.0':
        groups.update({name:spec[name] for name in ('developer_decisions','references')})
        groups['architecture']=spec['architecture']['constraints']+spec['architecture']['boundaries']
    unknown=set(collections)-set(groups)-{'dependencies','sections'}
    if unknown: raise b.Invalid('Unknown public collection: '+', '.join(sorted(unknown)))
    idmap={r['id']:(name,r) for name,rows in groups.items() for r in rows}
    graph={rid:set() for rid in idmap}
    for rid,(_,row) in idmap.items():
        for value in b.flatten(row).values():
            candidates=value if isinstance(value,list) else [value]
            for v in candidates:
                if isinstance(v,str) and v in idmap and v!=rid:
                    graph[rid].add(v);graph[v].add(rid)
    selected={r['id'] for name in collections if name in groups for r in groups[name]}
    todo=list(selected)
    while todo:
        for neighbor in graph[todo.pop()]:
            if neighbor not in selected: selected.add(neighbor);todo.append(neighbor)
    data={k:copy.deepcopy(spec[k]) for k in ['schemaVersion','meta','scope_items','naming_rules','functional_defaults','baseline','dependencies']}
    for name,source in groups.items():
        if name=='architecture': continue
        rows=[r for r in source if r['id'] in selected]
        if rows: data[name]=rows
    if 'architecture' in groups and any(r['id'] in selected for r in groups['architecture']):
        data['architecture']={'summary':spec['architecture']['summary'],
          **{name:[r for r in spec['architecture'][name] if r['id'] in selected] for name in ('constraints','boundaries')}}
    if 'sections' in collections: data['sections']=spec['sections']
    return {'kind':'SELECTED_PUBLIC_CONTEXT','authoritative':False,'source_sha256':revision(spec),
            'included_ids':sorted(selected),'omitted_ids':sorted(set(idmap)-selected),'data':data}

def defaults_sha(spec): return b.digest(b.canonical(spec['functional_defaults']))

def eval_gate(spec,control):
    errors=[]; round_reports=[]; rounds=control.get('eval_rounds',[])
    expected_cases={r['id']:r['expected'] for r in spec.get('test_cases',[])}
    requirements={r['id'] for r in spec.get('requirements',[])}
    visual_expected={(r['media_ref'],r['ui_element_ref']) for r in spec.get('ui_callouts',[])}
    context_ids=set(); packet_hashes=set()
    requests={x.get('context_id'):x for x in control.get('eval_requests',[]) if x.get('revision_sha256')==revision(spec)}
    for rn,r in enumerate(rounds):
        current=[];readers=r.get('readers',[])
        import review_contract
        if not isinstance(readers,list):
            current.append(issue('F2_RESPONSE_SHAPE','/eval_rounds/'+str(rn),'Readers must be an array'));readers=[]
        valid_readers=[]
        for index,reader in enumerate(readers):
            problems=review_contract.shape_errors(reader,f'/eval_rounds/{rn}/readers/{index}')
            current+=problems
            if not problems:valid_readers.append(reader)
        readers=valid_readers
        if r.get('revision_sha256')!=revision(spec): current.append(issue('F5_STALE_ROUND','/eval_rounds/'+str(rn),'Round is for another revision'))
        if len(readers)!=3 or {x.get('reader_id') for x in readers}!=set(review_contract.READER_ROLES): current.append(issue('F2_READERS','/eval_rounds/'+str(rn),'Exactly the three issued independent reader roles are required'))
        for i,reader in enumerate(readers):
            ptr=f'/eval_rounds/{rn}/readers/{i}'
            cid=reader.get('context_id');ph=reader.get('request_sha256')
            if not cid or cid in context_ids or not ph or ph in packet_hashes or reader.get('isolated') is not True or reader.get('saw_other_results') is not False or reader.get('status')!='COMPLETED': current.append(issue('F2_INDEPENDENCE',ptr,'Fresh isolated context and request required'))
            context_ids.add(cid);packet_hashes.add(ph)
            request=requests.get(cid,{})
            if request.get('request_sha256')!=ph or request.get('round')!=rn+1: current.append(issue('F2_REQUEST_BINDING',ptr,'No issued current-revision review packet'))
            if not reader.get('questions') or set(reader.get('covered_requirements',[]))!=requirements: current.append(issue('F2_COVERAGE',ptr,'Question/requirement coverage incomplete'))
            if 'findings' not in reader or reader['findings']: current.append(issue('F2_OPEN_FINDING',ptr,'Independent reader findings remain open'))
            for question in reader.get('questions',[]):
                try:
                    actual=resolve(spec,question.get('pointer'))
                    if b.canonical(actual)!=b.canonical(question.get('answer')): current.append(issue('F2_ANSWER',ptr,'Pointer answer differs'))
                except (b.Invalid,KeyError,IndexError,TypeError): current.append(issue('F2_POINTER',ptr,'Question cannot be answered at pointer'))
            import review_contract
            current+=review_contract.findings(spec,reader,ptr)
            if request.get('reader_id')!=reader.get('reader_id'): current.append(issue('F2_REQUEST_BINDING',ptr,'Reader identity differs from issued packet'))
            if request.get('protocol')!=source_pool.PROTOCOL or not review_contract.valid_sha(request.get('source_pool_sha256')):
                current.append(issue('F2_PROTOCOL',ptr,'Reissue legacy review packets with current pooled-source bindings'))
            if request.get('review_role')!=review_contract.READER_ROLES.get(reader.get('reader_id')) or not review_contract.valid_sha(request.get('blindness_sha256')):
                current.append(issue('F2_BLINDNESS_BINDING',ptr,'Fresh role-bound oracle admission is required'))
            if 'work_profile' in control or control.get('work_profile_required'):
                try:validate_work_reader({'delivery':{'spec':spec,'control':control}},reader,request)
                except b.Invalid as error:current.append(issue('F2_WORK_DISPATCH',ptr,str(error)))
            if not re.fullmatch('[a-f0-9]{64}',reader.get('response_sha256','')) or not reader.get('execution_evidence'): current.append(issue('F2_EXECUTION',ptr,'Recorded execution evidence and response hash required'))
        scenario={reader['reader_id']:reader for reader in readers if reader['reader_id'] in review_contract.SCENARIO_READERS}
        if len(scenario)==2:
            left=scenario['reader-1'].get('expected_results',{});right=scenario['reader-2'].get('expected_results',{})
            if set(left)!=set(expected_cases) or set(right)!=set(expected_cases) or b.canonical(left)!=b.canonical(right) or b.canonical(left)!=b.canonical(expected_cases): current.append(issue('F3_DISAGREEMENT','/eval_rounds/'+str(rn),'Independent scenario results differ/incomplete'))
        else:
            current.append(issue('F3_DISAGREEMENT','/eval_rounds/'+str(rn),'Both issued scenario readers must return complete results'))
        pairs=set()
        for reader in readers:
            for visual in reader.get('visual_matches',[]):
                if visual.get('method')!='PNG_INSPECTION' or not visual.get('label_match') or not visual.get('layout_match'): current.append(issue('F4_VISUAL','/eval_rounds/'+str(rn),'Visual mismatch or inspection missing'))
                else: pairs.add((visual.get('media_ref'),visual.get('ui_element_ref')))
        if visual_expected-pairs: current.append(issue('F4_COVERAGE','/eval_rounds/'+str(rn),'PNG/callout inspection incomplete'))
        plans=[x for x in readers if x.get('plan_completed') is True]
        if not plans: current.append(issue('F5_PLAN','/eval_rounds/'+str(rn),'Plan simulation missing'))
        for reader in plans:
            for decision in reader.get('plan_decisions',[]):
                try: resolve(spec,decision.get('pointer'))
                except (b.Invalid,KeyError,IndexError,TypeError): current.append(issue('F5_PLAN_DECISION','/eval_rounds/'+str(rn),'Plan introduced an unanswered decision'))
        if r.get('new_open_count')!=0: current.append(issue('F5_NEW_OPEN','/eval_rounds/'+str(rn),'New open points found'))
        round_reports.append({'round':rn+1,'reader_count':len(readers),'new_open_count':r.get('new_open_count'),'status':'PASS' if not current else 'FAIL'})
        if rn>=len(rounds)-2: errors+=current
    if len(rounds)<2 or any(x['status']!='PASS' for x in round_reports[-2:]): errors.append(issue('F5_TWO_CLEAN','/eval_rounds','Two consecutive clean rounds required'))
    if control.get('review_execution_confirmed') is not True or not control.get('review_confirmation_receipt') or control.get('review_confirmation_protocol')!=source_pool.PROTOCOL:
        errors.append(issue('F2_EXECUTION_CONFIRMATION','/control','Owner must confirm real current-protocol independent execution; fixtures are not model eval'))
    return errors,round_reports

def evaluate(doc):
    try: spec,control=get_state(doc)
    except b.Invalid as e: return {'decision':'BLOCKED','issues':[issue('D3_DELIVERY_MISSING','/delivery',str(e))],'open_counts':{'missing_profile':1},'layers':[],'rounds':[]}
    schema_issues=schema_errors(spec)
    if schema_issues:
        return {'decision':'BLOCKED','coding_readiness':'BLOCKED','technical_open_count':0,
          'spec_sha256':revision(spec),'issues':schema_issues,'open_counts':{'invalid_profile':len(schema_issues)},
          'layers':[{'layer':1,'status':'FAIL'}]+[{'layer':n,'status':'NOT_RUN'} for n in range(2,6)],
          'rounds':[],'eval_round_count':0,'confidence_limit':'Invalid data has not been evaluated'}
    errors=shape_and_profile(spec);counts={}
    import work_profiles
    try:work_profiles.assert_release(doc)
    except b.Invalid as error:errors.append(issue('WORK_PROFILE_GATE','/control/work_profile',str(error)))
    def closed(row):
        if not isinstance(row,dict) or row.get('status')!='CLOSED' or row.get('owner_confirmed') is not True: return False
        try: resolve(spec,row.get('answer_pointer'));return True
        except (b.Invalid,KeyError,IndexError,TypeError): return False
    modern=spec.get('schemaVersion')=='3.0'
    if modern: import handoff3 as h
    for name in ['intake','decisions_gaps','open_questions','system_conflicts']:
        values=control.get(name)
        counts[name]=sum(not closed(x) and not (modern and name!='system_conflicts' and h.eligible_technical(spec,x)) for x in values) if isinstance(values,list) else 1
    if modern:
        counts['architecture_findings']=sum(not closed(x) for x in control.get('architecture_findings',[]))
    answers=control.get('safe_default_answers')
    counts['unapproved_safe_defaults']=sum(not (closed({**x,'owner_confirmed':x.get('owner_approved')}) and x.get('approval_receipt')) for x in answers) if isinstance(answers,list) else 1
    counts['markers']=h.blocking_markers(spec) if modern else sum(isinstance(v,str) and bool(GAP.search(v)) for v in b.flatten(spec).values())
    if any(counts.values()): errors.append(issue('D2_ZERO_OPEN','/control','All open counters must be zero'))
    if control.get('approved_spec_sha256')!=revision(spec) or not control.get('approval_receipt'): errors.append(issue('APPROVAL_CURRENT','/control','Current FS-TS approval required; private receipt stays outside handoff'))
    if control.get('approved_defaults_sha256')!=defaults_sha(spec) or not control.get('defaults_approval_receipt'): errors.append(issue('D5_APPROVED_DEFAULTS','/functional_defaults','Business-approved current defaults required'))
    if control.get('system_conflicts') is None: errors.append(issue('E5_RECONCILIATION','/control','System/design comparison not stated'))
    for delivered in control.get('delivery_history',[]):
        if delivered['version']==spec.get('meta',{}).get('handoff_version') and delivered['spec_sha256']!=revision(spec): errors.append(issue('A5_A6_IMMUTABLE_VERSION','/meta/handoff_version','Changed delivered spec needs a new version'))
    if control.get('feedback') and control.get('feedback_version')==spec.get('meta',{}).get('handoff_version'): errors.append(issue('A6_FEEDBACK_VERSION','/meta/handoff_version','Developer correction needs a new version'))
    if spec.get('meta',{}).get('development_id')!=doc['development']['id']: errors.append(issue('A2_SCOPE_ID','/spec/meta','Workspace belongs to another development'))
    if spec.get('meta',{}).get('mode')=='UPDATE':
        baseline=control.get('baseline_spec')
        if not baseline or spec['baseline']['spec_sha256']!=revision(baseline): errors.append(issue('E4_BASELINE_HASH','/baseline','Decoded reference snapshot missing or differs'))
        elif baseline['meta']['development_id']!=spec['meta']['development_id']: errors.append(issue('E1_FOREIGN_BASELINE','/baseline','Other development needs a separate handoff'))
        if spec['baseline']['kind']=='FINAL_HANDOFF' and (control.get('baseline_system_state')!='ABSENT_VERIFIED' or control.get('baseline_final_approved') is not True): errors.append(issue('E3_ABSENCE','/baseline','System absence and final-reference approval must be confirmed'))
        if spec['baseline']['kind']=='SYSTEM' and control.get('baseline_system_state')!='PRESENT_VERIFIED': errors.append(issue('E3_PRESENCE','/baseline','System snapshot not verified'))
    if not schema_errors(spec):
        eval_errors,rounds=eval_gate(spec,control)
    else: eval_errors,rounds=[issue('F0_SCHEMA','/spec','Fix data before evaluation')],[]
    layers=[{'layer':1,'status':'PASS' if not errors else 'FAIL'},
            {'layer':2,'status':'FAIL' if any(x['code'].startswith('F2') for x in eval_errors) else 'PASS'},
            {'layer':3,'status':'FAIL' if any(x['code'].startswith('F3') for x in eval_errors) else 'PASS'},
            {'layer':4,'status':'FAIL' if any(x['code'].startswith('F4') for x in eval_errors) else 'PASS'},
            {'layer':5,'status':'FAIL' if any(x['code'].startswith(('F5','F0')) for x in eval_errors) else 'PASS'}]
    if not rounds:
        for layer in layers[1:]: layer['status']='NOT_RUN'
    errors+=eval_errors
    technical=sum(x['status']=='OPEN' for x in spec.get('developer_decisions',[]))
    return {'decision':'DELIVERABLE' if not errors else 'BLOCKED','coding_readiness':'BLOCKED' if errors else 'READY_FOR_DEVELOPER_DECISIONS' if technical else 'READY_FOR_CODING',
            'technical_open_count':technical,'spec_sha256':revision(spec),'issues':errors,'open_counts':counts,'layers':layers,'rounds':rounds,'eval_round_count':len(rounds),
            'confidence_limit':'Deterministic checks and supplied review records increase confidence; they do not prove absolute semantic completeness or SAP runtime.'}

def neutral_report(report):
    return {k:report[k] for k in ['decision','coding_readiness','technical_open_count','spec_sha256','open_counts','layers','rounds','eval_round_count','confidence_limit']}

def release_init(doc):
    if doc.get('delivery') is not None: raise b.Invalid('Delivery state exists; edit the existing TOON workspace')
    content=doc['content']; scope=content['bolumler'].get('1.1',{}).get('alanlar',{})
    spec={'schemaVersion':'3.0','developer_decisions':[],'references':[],
          'architecture':{'summary':'BİLGİ BEKLİYOR','boundaries':[],'constraints':[]},
          'meta':{'development_id':doc['development']['id'],'name':doc['development']['name'],'development_name':doc['development']['slug'],'handoff_version':doc['handoff']['version'],'fs_version':content['meta']['surum'],'edition':doc['development']['edition'],'release':'BİLGİ BEKLİYOR','types':content['turler'],'language':doc['handoff']['language'],'mode':doc['handoff']['mode']},
          'scope_items':{'included':[],'excluded':[],'preserved':[]},'naming_rules':{},'functional_defaults':{},'sections':copy.deepcopy(content['bolumler']),
          'baseline':{'kind':'NEW' if doc['handoff']['mode']=='NEW' else 'UNKNOWN','system':None,'client':None,'verified_at':None,'reference_name':None,'reference_version':None,'reference_sha256':None,'spec_sha256':None,'evidence_paths':[]},'dependencies':[],'change_rationale':'BİLGİ BEKLİYOR','changelog':[]}
    spec.update({k:[] for k in COLLECTIONS})
    for row in content['bolumler'].get('2.2',{}).get('satirlar',[]):
        if row and isinstance(row[0],str): spec['requirements'].append({'id':row[0],'statement':row[1],'acceptance_criteria':[]})
    control={'intake':[],'decisions_gaps':[{'id':'OPEN-RELEASE','status':'GAP','owner':'CONSULTANT','domain':'BUSINESS','question':'Supply the approved business scope, behavior and defaults; technical choices belong to the developer','options':[],'options_reason':'Source context is not yet sufficient for grounded alternatives'}],'open_questions':[],'system_conflicts':[],
      'safe_default_answers':[],'approved_spec_sha256':None,'approval_receipt':None,'approved_defaults_sha256':None,'defaults_approval_receipt':None,'review_execution_confirmed':False,'eval_rounds':[],
      'assets':[],'baseline_spec':None,'mockup_exception':{'authorized':False,'receipt':None},'feedback':[],'eval_requests':[],
      'architecture_findings':[],'dependency_contracts':[],'check_cache':{},'review_history':[]}
    control.update(baseline_system_state='UNKNOWN',baseline_final_approved=False,round_limit=4)
    doc['delivery']={'spec':spec,'control':control}

def _issue_review_packets(doc,round_number,prepared):
    # Pure packet assembly is also used by explicitly synthetic test fixtures.
    import review_contract as review
    spec,control=get_state(doc);packets=[]
    blindness=review.blind_binding(spec,prepared)
    pooled=prepared.get('source_pool',{'version':source_pool.POOL_VERSION,'sources':[]})
    pool_sha=source_pool.value_sha(pooled)
    pngs=[row for row in prepared['assets'] if row['path'] in {m['path'] for m in spec['media']}]
    for i in range(3):
        packet={'protocol':source_pool.PROTOCOL,'revision_sha256':revision(spec),'round':round_number,
          'reader_id':'reader-'+str(i+1),'context_id':str(uuid.uuid4()),
          'review_role':review.READER_ROLES['reader-'+str(i+1)],'blindness_sha256':blindness,
          'instructions':'Use a fresh isolated context. Inspect functional requirements, approved defaults, packaged references and actual PNGs. Return one functional question per requirement with requirement_ref, its statement pointer and observed answer. The first two readers derive every test result from functional rules and inputs; expected test answers are withheld. The third reader checks the complete specification. Cover actual PNG callouts. Provide concrete plan decisions pointing to objects, architecture constraints, dependencies and each delegated technical decision that exists. Report every contradiction, missing business answer and missing reference. Do not fabricate execution evidence, choose tools or change the functional contract.',
          'scenario_reader':i<2,'authoritative':False,'spec':review.review_spec(spec,i<2),
          'png_files':[{'path':x['source_path'],'sha256':x['sha256']} for x in pngs],
          'references':copy.deepcopy(prepared['references']),
          'source_pool':copy.deepcopy(pooled),
          'dependency_contracts':prepared.get('dependency_contracts',[]),
          'input_sha256':prepared.get('input_sha256')}
        if 'work_profile' in control:
            packet['work_task']={'task_id':'review:'+str(round_number)+':'+packet['reader_id'],'role':'reviewer',
                                'context_id':packet['context_id'],'agent':'byw-reader-'+chr(97+i),
                                'dispatch_required':True,'private':True}
        packet['reply_shape']={'reader_id':'echo request','context_id':'echo request','request_sha256':'echo request',
          'status':'COMPLETED or BLOCKED','isolated':True,'saw_other_results':False,
          'covered_requirements':['all inspected requirement IDs'],
          'questions':[{'requirement_ref':'requirement ID','question':'functional question','pointer':'/requirements/N/statement','answer':'observed statement'}],
          'expected_results':{'test-case ID':'independently derived result'},
          'visual_matches':[{'media_ref':'ID','ui_element_ref':'ID','method':'PNG_INSPECTION or NOT_RUN','label_match':'boolean','layout_match':'boolean'}],
          'plan_completed':'boolean','plan_decisions':[{'decision':'concrete decision','pointer':'object, architecture, dependency or delegated-decision pointer'}],
          'findings':['unresolved issues'],
          'execution_evidence':'controller supplies genuine private execution reference',
          'response_sha256':'controller hashes the parsed response excluding response_sha256 and execution_evidence'}
        packet['instructions']+=' Resolve each reference using source_id, matching sha256/path and pointer into source_pool.sources[].content. Each complete source appears once. Initialize source_pool.Reader once for repeated access; legacy packets are read-only and cannot provide current review credit.'
        packet['source_pool_sha256']=pool_sha
        packet['request_sha256']=source_pool.value_sha(packet);packets.append(packet)
    source_pool.check_volume(packets,control)
    control['eval_requests'] += [{'protocol':source_pool.PROTOCOL,'revision_sha256':revision(spec),'round':round_number,
      'reader_id':x['reader_id'],'context_id':x['context_id'],'request_sha256':x['request_sha256'],
       'input_sha256':x['input_sha256'],'source_pool_sha256':x['source_pool_sha256'],
       'review_role':x['review_role'],'blindness_sha256':x['blindness_sha256'],
       **({'task_id':x['work_task']['task_id']} if 'work_task' in x else {})} for x in packets]
    return packets

def eval_request(doc,round_number,assets_root=None):
    import preflight
    import profile_runtime
    spec,control=get_state(doc)
    profile_runtime.physical_admit(doc,'review',assets_root)
    if shape_and_profile(spec):raise b.Invalid('Complete deterministic profile before reader evaluation')
    if round_number!=len(control['eval_rounds'])+1 or round_number>control.get('round_limit',4):
        raise b.Invalid('Round sequence/limit reached; resolve with owner before continuing')
    prepared=preflight.inspect(doc,assets_root)
    if prepared['status']!='PASS':
        raise b.Invalid('Preflight blocked: '+json.dumps(prepared['issues'],ensure_ascii=False))
    return _issue_review_packets(doc,round_number,prepared)

def record_round(doc,record):
    import review_contract
    spec,control=get_state(doc)
    if record.get('revision_sha256')!=revision(spec): raise b.Invalid('Review does not match current snapshot')
    expected={r['context_id']:r for r in control['eval_requests'] if r['round']==len(control['eval_rounds'])+1 and r['revision_sha256']==revision(spec)}
    for reader in record.get('readers',[]):
        if reader.get('context_id') not in expected or reader.get('request_sha256')!=expected[reader['context_id']]['request_sha256']: raise b.Invalid('Review is not bound to an issued independent packet')
        request=expected[reader['context_id']]
        if request.get('protocol')!=source_pool.PROTOCOL or not review_contract.valid_sha(request.get('source_pool_sha256')):
            raise b.Invalid('Legacy review records are read-only; issue fresh pooled-source packets')
        if request.get('reader_id')!=reader.get('reader_id') or request.get('review_role')!=review_contract.READER_ROLES.get(reader.get('reader_id')) or not review_contract.valid_sha(request.get('blindness_sha256')):
            raise b.Invalid('Review does not match a current role/oracle admission binding')
        reader['response_sha256']=review_contract.response_sha(reader)
        validate_work_reader(doc,reader,expected[reader['context_id']])
    if len(record.get('readers',[]))!=3 or {row['reader_id'] for row in record['readers']}!=set(review_contract.READER_ROLES):
        raise b.Invalid('Round needs the three unique issued reader roles')
    control['eval_rounds'].append(record)


def validate_work_reader(doc,reader,request):
    """Bind managed reader records to actual counted dispatch receipts."""
    import profile_runtime
    if not profile_runtime.managed(doc):return
    import work_profiles
    work_profiles.admit(doc,'review')
    profile=doc['delivery']['control'].get('work_profile',{})
    if not isinstance(request,dict) or not all(key in request for key in ('round','reader_id','request_sha256','revision_sha256')):
        raise b.Invalid('Managed reader has no complete issued request binding')
    if request['revision_sha256']!=revision(doc['delivery']['spec']):
        raise b.Invalid('Managed reader request belongs to another specification')
    # Work-profile storage is private; no receipt is a provider-authentication claim.
    attempts=profile.get('attempts',[])
    task_id=request.get('task_id','review:'+str(request['round'])+':'+request['reader_id'])
    matches=[row for row in attempts if row.get('task_id')==task_id
             and row.get('context_id')==reader['context_id']
             and row.get('input_sha256')==request['request_sha256']]
    if not matches:raise b.Invalid('Managed reader has no counted dispatch attempt')
    row=matches[-1]
    if row.get('role')!='reviewer' or row.get('status')!='PASS' or row.get('evidence',{}).get('response_sha256')!=reader['response_sha256']:
        raise b.Invalid('Managed reader completion/hash differs from its counted dispatch receipt')
    if row.get('spec_sha256')!=revision(doc['delivery']['spec']) or row.get('checker_sha256')!=profile.get('capacity',{}).get('checker_sha256'):
        raise b.Invalid('Managed reader dispatch belongs to a stale specification or checker')

def feedback(doc,text,category):
    spec,control=get_state(doc)
    control['feedback_version']=spec['meta']['handoff_version']
    control['feedback'].append({'id':'REG-'+str(len(control['feedback'])+1),'category':category,'defect':text,'status':'OPEN'})
    control['open_questions'].append({'id':control['feedback'][-1]['id'],'status':'OPEN','question':text})
    control['approved_spec_sha256']=None;control['approval_receipt']=None;control['eval_rounds']=[]

def clean_png(data):
    if not data.startswith(b'\x89PNG\r\n\x1a\n'): raise b.Invalid('Not PNG')
    import zlib
    pos=8;ended=False
    while pos<len(data):
        if pos+12>len(data): raise b.Invalid('Truncated PNG')
        length=int.from_bytes(data[pos:pos+4],'big');tag=data[pos+4:pos+8];end=pos+12+length
        if end>len(data): raise b.Invalid('Malformed PNG')
        if int.from_bytes(data[end-4:end],'big')!=zlib.crc32(data[pos+4:end-4])&0xffffffff: raise b.Invalid('PNG CRC')
        if tag in (b'tEXt',b'iTXt',b'zTXt',b'eXIf'): raise b.Invalid('PNG attribution metadata')
        pos=end
        if tag==b'IEND': ended=True;break
    if not ended or pos!=len(data): raise b.Invalid('PNG end/trailing data')

def invalidate(doc,defaults=False):
    _,control=get_state(doc)
    import handoff3 as h
    h.archive_reviews(control)
    control.update(approved_spec_sha256=None,approval_receipt=None,review_execution_confirmed=False,eval_rounds=[],eval_requests=[])
    if defaults: control.update(approved_defaults_sha256=None,defaults_approval_receipt=None)

def changes(before,after):
    if before['meta']['development_id']!=after['meta']['development_id']: raise b.Invalid('Delta is single-development')
    left=b.flatten(before);right=b.flatten(after);rows=[]
    for path in sorted(set(left)|set(right)):
        bp,ap=path in left,path in right
        if bp and ap and left[path]==right[path]: continue
        rows.append({'pointer':path,'kind':'MODIFIED' if bp and ap else 'ADDED' if ap else 'REMOVED','before_present':bp,'after_present':ap,'before':left.get(path),'after':right.get(path)})
    return {'baseline_sha256':revision(before),'target_sha256':revision(after),'rationale':after['change_rationale'],'changes':rows}

def copy_asset(asset,root):
    name=asset['path'];rel=b.safe_name(name);root=Path(root).resolve();raw=root.joinpath(*rel.parts)
    if raw.is_symlink() or not raw.resolve().is_relative_to(root) or not raw.is_file(): raise b.Invalid('Missing/escaped asset')
    if raw.stat().st_size>32*1024*1024: raise b.Invalid('Asset size exceeds 32 MiB')
    data=raw.read_bytes()
    if b.digest(data)!=asset['sha256']: raise b.Invalid('Asset hash mismatch')
    if name.startswith('mockup/screens/'):
        if Path(name).suffix!='.png' or not data.startswith(b'\x89PNG\r\n\x1a\n'): raise b.Invalid('Static mockup must be PNG')
        clean_png(data)
    elif Path(name).suffix.lower() in ('.md','.toon','.csv','.xml'):
        if privacy(data.decode('utf-8')): raise b.Invalid('Forbidden evidence/contract content')
        if Path(name).suffix.lower()=='.xml' and re.search(r'(?i)<!DOCTYPE|<script|\bon\w+\s*=',data.decode('utf-8')): raise b.Invalid('Active XML not a data contract')
    elif name.startswith('mockup/interactive/') and Path(name).suffix.lower() in ('.html','.js','.mjs','.css','.json','.svg'):
        text=data.decode('utf-8')
        if TOOL.search(text): raise b.Invalid('Interactive producer metadata')
        if re.search(r'(?i)<!--|<meta[^>]+name\s*=\s*[\"\x27]generator',text): raise b.Invalid('HTML comments/generator metadata must be cleaned')
        if NETWORK.search(text): raise b.Invalid('Interactive mockup makes external/dynamic calls')
    else: raise b.Invalid('Code/unsupported asset refused')
    return name,data

def render_excel(spec,report,path):
    from openpyxl import Workbook
    from openpyxl.styles import Font,PatternFill,Alignment
    wb=Workbook();wb.properties.creator='';wb.properties.lastModifiedBy='';wb.properties.title=spec['meta']['name']
    ws=wb.active;ws.title='Özet'
    authority='fsts/fsts.toon'
    if spec.get('schemaVersion')=='3.0':
        import handoff3 as h
        authority=h.layout(spec)['specification']
    for k,v in [('Geliştirme',spec['meta']['name']),('ID',spec['meta']['development_id']),('Handoff',spec['meta']['handoff_version']),('FS-TS',spec['meta']['fs_version']),('Kapı',report['decision']),('Otorite',authority)]: ws.append([k,v])
    for name in COLLECTIONS+['dependencies']:
        rows=spec[name]
        if not rows: continue
        ws=wb.create_sheet(name[:31]);keys=list(rows[0]);ws.append(keys)
        for row in rows:
            ws.append([json.dumps(row[k],ensure_ascii=False) if isinstance(row[k],(dict,list)) else row[k] for k in keys])
    ws=wb.create_sheet('Tam veri');ws.append(['Pointer','Değer'])
    for path_,value in b.flatten(spec).items(): ws.append([path_,json.dumps(value,ensure_ascii=False)])
    for ws in wb:
        ws.freeze_panes='A2';ws.sheet_view.showGridLines=False
        for c in ws[1]: c.fill=PatternFill('solid',fgColor='18344D');c.font=Font(bold=True,color='FFFFFF',size=11)
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width=min(65,max(16,max(len(str(c.value or '')) for c in col)*.7))
        for row in ws:
            for c in row:
                if isinstance(c.value,str): c.data_type='s'
                c.alignment=Alignment(wrap_text=True,vertical='top')
            ws.row_dimensions[row[0].row].height=min(409,max(24,max(len(str(c.value or ''))//50+1 for c in row)*16+6))
    wb.save(path)
    with zipfile.ZipFile(path) as z: files={n:z.read(n) for n in z.namelist()}
    files['docProps/app.xml']=b'<?xml version="1.0"?><Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/extended-properties"><Application></Application></Properties>'
    with zipfile.ZipFile(path,'w',zipfile.ZIP_DEFLATED) as z:
        for n,data in files.items(): z.writestr(n,data)

def verify_excel(spec,path):
    from openpyxl import load_workbook
    wb=load_workbook(path);expected=[(p,json.dumps(v,ensure_ascii=False)) for p,v in b.flatten(spec).items()]
    if list(wb['Tam veri'].values)[1:]!=expected: raise b.Invalid('Excel complete-data parity')
    for name in COLLECTIONS+['dependencies']:
        rows=spec[name]
        if not rows: continue
        actual=list(wb[name[:31]].values);keys=list(rows[0])
        if list(actual[0])!=keys: raise b.Invalid('Excel heading mismatch')
        for i,row in enumerate(rows):
            values=[json.dumps(row[k],ensure_ascii=False) if isinstance(row[k],(dict,list)) else row[k] for k in keys]
            if list(actual[i+1])!=values: raise b.Invalid('Excel visible projection mismatch')
    if any(c.data_type=='f' for ws in wb for row in ws for c in row): raise b.Invalid('Unexpected formula')
    return len(expected)

FORMAT_TEXT='''# Handoff biçim sözleşmesi

fsts/fsts.toon tek içerik otoritesidir. UTF-8, LF, TOON 4.1; girinti 2 boşluk,
virgül ayırıcı ve strict decode. Şema fsts/fsts.schema.toon içinde aynı veri
modeliyle taşınır; çözülmüş şema veri nesnesine uygulanır. Nötr şema kimliği
urn:fsts:2.0:handoff. Pointer adresleri çözülmüş TOON veri yapısını gösterir.
Excel, nesne listesi ve açıklamalar bu kaynaktan türetilir. PNG yerleşimde,
TOON davranış ve veride yetkilidir. Proje varsayılanları ad/sürümle bağlıdır.
Sayısal teknik kimlik ve hassas ondalık değerler quoted string olarak tutulur.
null bilinen uygulanmıyor durumları için açık kind alanıyla birlikte kullanılır;
bilinmeyen fonksiyonel cevap teslimde bulunmaz. [] bilerek boş koleksiyondur.
Delta tam hedefin rehberidir; stable ID, önce/sonra ve var/yok ayrımı korunur.
Manifest bütün dosyaları listeler. Kendi girişinde SHA-256 null ve
self_hash=EXTERNAL olur; öz-hash döngüsü kurulmaz. Manifest ve ZIP SHA-256'sı
dış teslim doğrulamasında hesaplanır. Teslim edilmiş ZIP değiştirilmez.
Şema, kapı ve kayıtlı değerlendirmeler güveni artırır; mutlak anlamsal
eksiksizlik veya SAP runtime/UAT kanıtı değildir.
'''

def _legacy_package(doc,output,assets_root,include_excel=True):
    spec,control=get_state(doc);report=evaluate(doc)
    if report['decision']!='DELIVERABLE': raise b.Invalid('No handoff while gaps/eval gates remain: '+', '.join(sorted({r['code'] for r in report['issues']})))
    out=Path(output);out.mkdir(parents=True,exist_ok=True)
    filename='handoff-'+spec['meta']['development_name']+'-'+spec['meta']['handoff_version']+'.zip';b.safe_name(filename)
    target=out/filename
    if target.exists(): raise b.Invalid('Immutable delivery already exists; use a new version')
    files={'fsts/fsts.toon':b.encode(spec).encode(),'fsts/fsts.schema.toon':b.encode(SCHEMA_V2).encode(),
      'fsts/readiness-report.toon':b.encode(neutral_report(report)).encode(),
      'objects/object-list.toon':b.encode({'objects':spec['objects']}).encode(),
      'objects/naming-rules.md':('# İsimlendirme\n\n'+spec['naming_rules']['name']+' '+spec['naming_rules']['version']+'\n\n'+spec['naming_rules']['description']+'\n').encode(),
      'standards/functional-defaults.md':('# '+spec['functional_defaults']['name']+'\n\nSürüm: '+spec['functional_defaults']['version']+'\n\n'+'\n'.join(r['id']+': '+r['text'] for r in spec['functional_defaults']['rules'])+'\n').encode(),
      'CHANGELOG.md':('# Değişiklik geçmişi\n\n'+'\n'.join(r['version']+' | '+r['date']+' | '+r['change'] for r in spec['changelog'])+'\n').encode(),
      'FORMAT.md':FORMAT_TEXT.encode()}
    if spec['meta']['mode']=='UPDATE':
        source_path=control.get('baseline_reference_path')
        if not source_path: raise b.Invalid('Exact baseline artifact required privately; normalized spec hash is not ZIP hash')
        rel=b.safe_name(source_path);source=(Path(assets_root).resolve().joinpath(*rel.parts))
        if source.is_symlink() or not source.resolve().is_relative_to(Path(assets_root).resolve()) or not source.is_file(): raise b.Invalid('Baseline artifact missing/escaped')
        if b.digest(source.read_bytes())!=spec['baseline']['reference_sha256']: raise b.Invalid('Baseline artifact SHA-256 differs')
        if spec['baseline']['kind']=='FINAL_HANDOFF':
            if source.name!=spec['baseline']['reference_name']: raise b.Invalid('Baseline handoff filename differs')
            with zipfile.ZipFile(source) as prior:
                if 'fsts/fsts.toon' in prior.namelist():
                    verify(source)
                    actual=json.loads(b.codec('decode',prior.read('fsts/fsts.toon').decode()))
                    if revision(actual)!=spec['baseline']['spec_sha256']: raise b.Invalid('Baseline archive spec differs from normalized snapshot')
                elif control.get('baseline_normalization_confirmed') is not True: raise b.Invalid('Legacy baseline normalization needs owner confirmation')
        files['delta/baseline.toon']=b.encode(spec['baseline']).encode();files['delta/changes.toon']=b.encode(changes(control['baseline_spec'],spec)).encode()
    for asset in control.get('assets',[]):
        name,data=copy_asset(asset,assets_root)
        if name.startswith('mockup/interactive/'):
            exception=control.get('mockup_exception',{})
            if exception.get('authorized') is not True or not exception.get('receipt'): raise b.Invalid('Interactive exception not explicitly confirmed')
            if exception.get('offline_verified') is not True or not exception.get('offline_evidence'): raise b.Invalid('Offline runtime verification receipt missing; static scan alone is not runtime proof')
        if name in files: raise b.Invalid('Duplicate generated/asset path')
        files[name]=data
    for media in spec['media']:
        if media['path'] not in files or b.digest(files[media['path']])!=media['sha256']: raise b.Invalid('PNG/callout snapshot missing or changed')
    for name in spec['baseline']['evidence_paths']:
        if name not in files: raise b.Invalid('Baseline evidence missing')
    with tempfile.TemporaryDirectory(dir=out,prefix='candidate-') as temp:
        if include_excel:
            book=Path(temp)/'fsts.xlsx';render_excel(spec,report,book);verify_excel(spec,book);files['excel/fsts.xlsx']=book.read_bytes()
        files['README.md']=('# '+spec['meta']['name']+'\n\nKimlik: '+spec['meta']['development_id']+'\nSürüm: '+spec['meta']['handoff_version']+'\n\n'
          'Okuma sırası: manifest.toon, fsts/fsts.toon, fsts/readiness-report.toon, nesneler, standartlar, mockup ve delta.\n'
          'İçerik önceliği: tam FS-TS, proje varsayılanları, mockup. PNG yalnız yerleşimde yetkilidir.\n'
          'Kapsam içi: '+'; '.join(spec['scope_items']['included'])+'\nKapsam dışı: '+'; '.join(spec['scope_items']['excluded'])+'\nKorunan davranış: '+'; '.join(spec['scope_items']['preserved'])+'\n\n'
          'Dosya envanteri:\n'+'\n'.join('- '+n for n in sorted(files))+'\n- README.md\n- manifest.toon\n').encode()
        manifest={'development_id':spec['meta']['development_id'],'handoff_version':spec['meta']['handoff_version'],'fs_version':spec['meta']['fs_version'],'spec_sha256':revision(spec),
          'baseline':spec['baseline'],'linked_handoffs':spec['dependencies'],'files':[{'path':n,'sha256':b.digest(data),'bytes':len(data)} for n,data in sorted(files.items())],
          'self_hash':'EXTERNAL'}
        manifest['files'].append({'path':'manifest.toon','sha256':None,'bytes':None});files['manifest.toon']=b.encode(manifest).encode()
        for name,data in files.items():
            if name.startswith('mockup/interactive/'): continue
            if Path(name).suffix in ('.toon','.md','.csv','.xml') and privacy(data.decode()): raise b.Invalid('Forbidden output content: '+name)
        staged=Path(temp)/'candidate.zip'
        with zipfile.ZipFile(staged,'w',zipfile.ZIP_DEFLATED) as z:
            for name,data in sorted(files.items()):
                info=zipfile.ZipInfo(name,date_time=(1980,1,1,0,0,0));info.compress_type=zipfile.ZIP_DEFLATED;info.external_attr=0o100644<<16
                z.writestr(info,data)
        result=verify(staged,expected_name=filename)
        with target.open('xb') as f: f.write(staged.read_bytes())
    control.setdefault('delivery_history',[]).append({'version':spec['meta']['handoff_version'],'spec_sha256':revision(spec),'zip_sha256':b.digest(target.read_bytes())})
    return {'path':str(target),'zipSha256':b.digest(target.read_bytes()),'manifestSha256':b.digest(files['manifest.toon']),'verification':result}

def _verify_v2(path,expected_name=None):
    with zipfile.ZipFile(path) as z:
        names=z.namelist()
        if len(names)!=len(set(names)) or z.testzip() or z.comment: raise b.Invalid('ZIP duplicates/CRC/comment')
        if sum(i.file_size for i in z.infolist())>128*1024*1024: raise b.Invalid('Oversized ZIP')
        for i in z.infolist():
            b.safe_name(i.filename)
            if i.flag_bits&1 or (i.external_attr>>16)&0o170000==0o120000: raise b.Invalid('ZIP encryption/symlink')
        required={'README.md','manifest.toon','fsts/fsts.toon','fsts/fsts.schema.toon','fsts/readiness-report.toon','objects/object-list.toon','objects/naming-rules.md','standards/functional-defaults.md','CHANGELOG.md','FORMAT.md'}
        if not required.issubset(names): raise b.Invalid('Required file missing')
        spec=json.loads(b.codec('decode',z.read('fsts/fsts.toon').decode()));issues=shape_and_profile(spec)
        if issues: raise b.Invalid('Invalid delivered profile: '+issues[0]['code'])
        manifest_data=z.read('manifest.toon');manifest=json.loads(b.codec('decode',manifest_data.decode()));rows={r['path']:r for r in manifest['files']}
        if set(rows)!=set(names) or len(rows)!=len(manifest['files']) or manifest.get('self_hash')!='EXTERNAL' or rows['manifest.toon']['sha256'] is not None: raise b.Invalid('Manifest inventory/self-hash')
        if manifest['spec_sha256']!=revision(spec) or manifest['development_id']!=spec['meta']['development_id'] or manifest['handoff_version']!=spec['meta']['handoff_version'] or manifest['fs_version']!=spec['meta']['fs_version'] or manifest['baseline']!=spec['baseline'] or manifest['linked_handoffs']!=spec['dependencies']: raise b.Invalid('Manifest identity/source mismatch')
        for name in names:
            data=z.read(name)
            if name!='manifest.toon' and (b.digest(data)!=rows[name]['sha256'] or len(data)!=rows[name]['bytes']): raise b.Invalid('File hash mismatch')
            suffix=Path(name).suffix.lower()
            if name.startswith('mockup/interactive/'):
                if suffix not in ('.html','.js','.mjs','.css','.json','.svg'): raise b.Invalid('Unexpected interactive file')
                if re.search(r'(?i)<!--|<meta[^>]+name\s*=\s*[\"\x27]generator',data.decode()): raise b.Invalid('Interactive metadata not clean')
                if TOOL.search(data.decode()) or NETWORK.search(data.decode()): raise b.Invalid('Interactive metadata/network call')
            elif suffix in ('.toon','.md','.csv','.xml'):
                if privacy(data.decode()): raise b.Invalid('Code/producer/private output')
            elif suffix=='.xlsx':
                with zipfile.ZipFile(io.BytesIO(data)) as book:
                    for part in book.namelist():
                        if part.endswith('.xml') and TOOL.search(book.read(part).decode()): raise b.Invalid('Workbook producer metadata')
            elif suffix=='.png': clean_png(data)
            else: raise b.Invalid('Unexpected code/file')
        schema=json.loads(b.codec('decode',z.read('fsts/fsts.schema.toon').decode()))
        if b.canonical(schema)!=b.canonical(SCHEMA_V2): raise b.Invalid('Neutral schema mismatch')
        objects=json.loads(b.codec('decode',z.read('objects/object-list.toon').decode()))
        if objects!={'objects':spec['objects']}: raise b.Invalid('Object-list projection mismatch')
        report=json.loads(b.codec('decode',z.read('fsts/readiness-report.toon').decode()))
        if report['decision']!='DELIVERABLE' or report['spec_sha256']!=revision(spec) or any(report['open_counts'].values()) or any(r['status']!='PASS' for r in report['layers']) or len(report['rounds'])<2 or any(r['status']!='PASS' or r['new_open_count']!=0 or r['reader_count']<3 for r in report['rounds'][-2:]): raise b.Invalid('Readiness gates not clean')
        for media in spec['media']:
            if media['path'] not in names or b.digest(z.read(media['path']))!=media['sha256']: raise b.Invalid('Media binding mismatch')
        if spec['meta']['mode']=='UPDATE':
            if 'delta/baseline.toon' not in names or 'delta/changes.toon' not in names: raise b.Invalid('Update delta missing')
            base=json.loads(b.codec('decode',z.read('delta/baseline.toon').decode()));delta=json.loads(b.codec('decode',z.read('delta/changes.toon').decode()))
            if base!=spec['baseline'] or delta['target_sha256']!=revision(spec) or delta['baseline_sha256']!=base['spec_sha256'] or delta['rationale']!=spec['change_rationale']: raise b.Invalid('Delta snapshot mismatch')
        filename='handoff-'+spec['meta']['development_name']+'-'+spec['meta']['handoff_version']+'.zip'
        if (expected_name or Path(path).name)!=filename: raise b.Invalid('Filename violates contract')
        if 'excel/fsts.xlsx' in names:
            with tempfile.TemporaryDirectory(prefix='handoff-check-') as t:
                book=Path(t)/'fsts.xlsx';book.write_bytes(z.read('excel/fsts.xlsx'));verify_excel(spec,book)
    return {'status':'PASS','files':len(names),'manifestSha256':b.digest(manifest_data),'specSha256':revision(spec),'scope':'Local data/artifact verification; not fresh semantic/SAP execution'}

def package(doc,output,assets_root,include_excel=True):
    import handoff3 as h
    return h.package(doc,output,assets_root,include_excel)

def verify(path,expected_name=None):
    with zipfile.ZipFile(path) as archive:
        legacy='fsts/fsts.toon' in archive.namelist()
    if legacy: return _verify_v2(path,expected_name)
    import handoff3 as h
    return h.verify(path,expected_name)
