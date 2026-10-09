#!/usr/bin/env python3
# SPDX-License-Identifier: GPL-3.0-only
# Modified derivative release 2.0.0, 2026-09-30; see NOTICE.md.
"""Provider-neutral FS-TS TOON, projection and single-development handoff."""
import argparse, copy, hashlib, io, json, math, os, re, shutil, subprocess, sys, tempfile, zipfile
from pathlib import Path, PurePosixPath
from decimal import Decimal
from datetime import date
import legacy_core as legacy

ROOT = Path(__file__).resolve().parent.parent
ID = re.compile(r'^[A-Z][A-Z0-9_]*-\d{2,}$')
PRODUCER = re.compile(r'(?i)\b(?:codex|claude|anthropic|openai|openpyxl|jsonschema|gpt-[\w.-]+|token-cost-auditor)\b|SKILL\.md|CLAUDE_PLUGIN_ROOT|\.claude-plugin|\.codex-plugin|scripts[/\\]bv[\w.-]*')
IMPLEMENTATION = re.compile(r'(?im)^\s*(?:CLASS\s+\w+\s+DEFINITION|METHOD\s+\w+\.|define\s+(?:root\s+)?view\s+entity|function\s+\w+\s*\(|import\s+.+\s+from\s+[\'\"])|```(?:abap|javascript|typescript|python|powershell|bash)\b')
BAD_FIELDS = {'producer','producerMetadata','generator','llm','llmModel','modelProvider','usedSkills','usedPlugins','producerTools','systemPrompt','developerPrompt','implementationCode'}

class Invalid(ValueError): pass

def canonical(value):
    return json.dumps(value,ensure_ascii=False,sort_keys=True,separators=(',',':'),allow_nan=False).encode('utf-8')

def digest(data): return hashlib.sha256(data).hexdigest()

def strict_json(text):
    def pairs(items):
        result={}
        for k,v in items:
            if k in result: raise Invalid('Duplicate JSON key: '+k)
            result[k]=v
        return result
    def exact_float(raw):
        value=float(raw)
        if not math.isfinite(value) or Decimal(str(value))!=Decimal(raw): raise Invalid('Quote exact high-precision decimals before migration')
        return value
    return json.loads(text,object_pairs_hook=pairs,parse_float=exact_float,parse_constant=lambda x:(_ for _ in ()).throw(Invalid('Non-finite number')))

def snapshot_sha(doc):
    return digest(canonical({k:v for k,v in doc.items() if k!='approval'}))

def ensure_numbers(value):
    if isinstance(value,float) and not math.isfinite(value): raise Invalid('Non-finite number')
    if isinstance(value,int) and not isinstance(value,bool) and abs(value)>9007199254740991:
        raise Invalid('Large integer must be a quoted string')
    if isinstance(value,dict):
        for v in value.values(): ensure_numbers(v)
    if isinstance(value,list):
        for v in value: ensure_numbers(v)

def codec(mode, text):
    result=subprocess.run([os.environ.get('BY_NODE','node'),str(ROOT/'scripts/codec.mjs'),mode],input=text,
                          text=True,encoding='utf-8',capture_output=True,timeout=30,cwd=ROOT)
    if result.returncode: raise Invalid('TOON structure: '+result.stderr.strip())
    return result.stdout

def encode(value):
    ensure_numbers(value)
    encoded=codec('encode',json.dumps(value,ensure_ascii=False,allow_nan=False))
    if canonical(json.loads(codec('decode',encoded))) != canonical(value):
        raise Invalid('TOON roundtrip changed data or numeric precision; quote exact decimals')
    return encoded

def encode_many(values):
    ensure_numbers(values)
    rows=json.loads(codec('batch-encode',json.dumps(values,ensure_ascii=False,allow_nan=False)))
    for key,value in values.items():
        if canonical(rows[key]['decoded']) != canonical(value):
            raise Invalid('TOON roundtrip changed data or numeric precision; quote exact decimals')
    return {key:row['text'] for key,row in rows.items()}

def decode_many(values):
    return json.loads(codec('batch-decode',json.dumps(values,ensure_ascii=False,allow_nan=False)))

def load(path):
    p=Path(path)
    if p.stat().st_size>16*1024*1024: raise Invalid('Document exceeds 16 MiB')
    return json.loads(codec('decode',p.read_text(encoding='utf-8')))

def save(value,path,overwrite=False):
    p=Path(path)
    if p.exists() and not overwrite: raise Invalid('Target exists; use a new path/version')
    p.parent.mkdir(parents=True,exist_ok=True)
    text=encode(value)
    with tempfile.NamedTemporaryFile('w',encoding='utf-8',newline='\n',dir=p.parent,delete=False) as f:
        f.write(text); temp=Path(f.name)
    try: os.replace(temp,p)
    finally: temp.unlink(missing_ok=True)

def schema_check(doc):
    from jsonschema import Draft202012Validator
    schema=json.loads((ROOT/'schema/fs-ts.schema.json').read_text(encoding='utf-8'))
    errors=sorted(Draft202012Validator(schema).iter_errors(doc),key=lambda e:str(e.path))
    return [{'code':'SCHEMA','path':'/'+ '/'.join(map(str,e.path)),'message':e.message} for e in errors]

def neutrality(value,path=''):
    errors=[]
    if isinstance(value,dict):
        for k,v in value.items():
            if k in BAD_FIELDS: errors.append({'code':'PRODUCER_FIELD','path':path+'/'+k,'message':'Producer/model/implementation field is forbidden'})
            errors+=neutrality(v,path+'/'+k)
    elif isinstance(value,list):
        for i,v in enumerate(value): errors+=neutrality(v,path+'/'+str(i))
    elif isinstance(value,str):
        if PRODUCER.search(value): errors.append({'code':'PRODUCER_ATTRIBUTION','path':path,'message':'Producer/tool/model attribution is forbidden'})
        if IMPLEMENTATION.search(value): errors.append({'code':'IMPLEMENTATION','path':path,'message':'Implementation source is forbidden in specification'})
    return errors

def records(doc):
    result=[]
    for sec,body in doc['content'].get('bolumler',{}).items():
        columns=legacy.BOLUM.get(sec,{}).get('sutunlar',[])
        if not columns or not columns[0].get('kimlik'): continue
        for row in body.get('satirlar',[]) if isinstance(body,dict) else []:
            if isinstance(row,list) and row and isinstance(row[0],str) and ID.fullmatch(row[0]): result.append((row[0],sec,row))
    for sec,rows in doc['extensions'].items():
        for row in rows: result.append((row['id'],'extensions/'+sec,row))
    return result

def validate(doc,baseline=None):
    errors=schema_check(doc); blockers=[]; gaps=[]; checks=[]
    if errors: return {'structure':'FAIL','errors':errors,'gaps':[],'blockers':['SCHEMA'],'readiness':'BLOCKED','checks':[]}
    errors+=neutrality(doc)
    scope=doc['content']['bolumler'].get('1.1',{}).get('alanlar',{})
    if scope.get('gel_id')!=doc['development']['id']: errors.append({'code':'SINGLE_DEVELOPMENT','path':'/content/1.1','message':'Development identity mismatch'})
    if scope.get('sistem') not in ('—',doc['development']['edition']) and not legacy.isaretli(scope.get('sistem','')):
        errors.append({'code':'EDITION','path':'/development/edition','message':'Edition differs from FS-TS'})
    all_records=records(doc); ids=set()
    for rid,sec,row in all_records:
        if rid in ids: errors.append({'code':'DUPLICATE_ID','path':sec,'message':rid})
        ids.add(rid)
        if isinstance(row,dict):
            if row.get('developmentId',doc['development']['id'])!=doc['development']['id']:
                errors.append({'code':'FOREIGN_MUTATION','path':sec,'message':'Separate handoff required'})
    for rid,sec,row in all_records:
        if isinstance(row,dict):
            for ref in row.get('refs',[]):
                if ref not in ids: errors.append({'code':'UNRESOLVED_REF','path':sec+'/'+rid,'message':ref})
    project={'katalog':[c['name'] for c in doc['project']['catalog'] if c['status']=='VERIFIED']}
    try:
        findings,_,_=legacy.denetle(doc['content'],project,'')
        for f in findings:
            item={'code':f['kural'],'path':f['yol'],'message':f['mesaj']}
            if f['kural']=='UYD_001':
                gaps.append(item); blockers.append('SAP_OBJECT_PROVENANCE')
            elif f.get('acik') or f['sev']!='HATA': gaps.append(item)
            else: errors.append(item)
            if f['kural'] in ('ZINCIR_002','ZINCIR_003'): blockers.append('TRACEABILITY')
    except (KeyError,TypeError,IndexError) as e:
        errors.append({'code':'LEGACY_SHAPE','path':'/content','message':str(e)})
    for _,_,text in legacy.hucreler(doc['content']):
        if legacy.isaretli(text) or '[DOĞRULANACAK]' in text: blockers.append('UNRESOLVED_INPUT')
    for row in doc['content']['bolumler'].get('7.3',{}).get('satirlar',[]):
        if len(row)>5 and row[5]=='Açık': blockers.append('OPEN_POINT:'+row[0])
    for name in ('release','objects','authorization'):
        v=doc['verification'][name]
        if v['status'] not in ('VERIFIED','NOT_APPLICABLE'): blockers.append('VERIFICATION:'+name)
        if v['status']=='VERIFIED' and not v['evidenceRefs']: errors.append({'code':'MISSING_EVIDENCE','path':'/verification/'+name,'message':'VERIFIED needs scoped evidence'})
    if doc['development']['release'] is None or doc['development']['edition']=='UNKNOWN': blockers.append('TARGET_RELEASE')
    for dep in doc['dependencies']:
        if dep['developmentId']==doc['development']['id']: errors.append({'code':'SELF_DEPENDENCY','path':'/dependencies','message':'Dependency must reference another development'})
        if dep['version'] is None or dep['status'] not in ('VERIFIED','NOT_APPLICABLE'): blockers.append('DEPENDENCY:'+dep['developmentId'])
    ap=doc['approval']
    if ap['status']=='APPROVED_SNAPSHOT':
        if ap['payloadSha256']!=snapshot_sha(doc) or not all(ap[k] for k in ('by','date','receipt')):
            errors.append({'code':'APPROVAL_SNAPSHOT','path':'/approval','message':'Approval does not bind current payload'})
    else: blockers.append('APPROVAL')
    b=doc['baseline']
    if doc['handoff']['mode']=='UPDATE':
        if b['kind'] in ('NEW','UNKNOWN'): blockers.append('BASELINE_UNKNOWN')
        else:
            if b['kind']=='SYSTEM' and b['systemState']!='PRESENT_VERIFIED': errors.append({'code':'SYSTEM_BASELINE','path':'/baseline','message':'System presence must be verified'})
            if b['kind']=='APPROVED_HANDOFF' and b['systemState']!='ABSENT_VERIFIED': errors.append({'code':'ABSENCE_NOT_PROVEN','path':'/baseline','message':'No access is not verified absence'})
            if not b['evidenceRefs'] or not b['verifiedAt']: blockers.append('BASELINE_EVIDENCE')
            if baseline is None: blockers.append('BASELINE_NOT_SUPPLIED')
            else:
                if baseline['development']['id']!=doc['development']['id']: errors.append({'code':'BASELINE_ID','path':'/baseline','message':'Baseline belongs to another development'})
                if b['payloadSha256']!=snapshot_sha(baseline): errors.append({'code':'BASELINE_HASH','path':'/baseline','message':'Wrong reference snapshot'})
                if b['kind']=='APPROVED_HANDOFF' and (baseline['approval']['status']!='APPROVED_SNAPSHOT' or baseline['approval']['payloadSha256']!=snapshot_sha(baseline)):
                    errors.append({'code':'BASELINE_APPROVAL','path':'/baseline','message':'Reference design is not an approved snapshot'})
    elif b['kind']!='NEW': errors.append({'code':'NEW_BASELINE','path':'/baseline','message':'NEW mode requires NEW baseline'})
    if doc['handoff']['mockupCodeException']['authorized'] and not doc['handoff']['mockupCodeException']['receipt']:
        errors.append({'code':'MOCKUP_EXCEPTION','path':'/handoff','message':'Explicit exception receipt required'})
    if 'U' in doc['content']['turler'] and doc['mockup']['status']=='NOT_CREATED': blockers.append('MOCKUP')
    if doc['mockup']['status']=='REFERENCED' and (not doc['mockup']['url'] or not doc['mockup']['version']): errors.append({'code':'MOCKUP_REFERENCE','path':'/mockup','message':'Fixed URL and version required'})
    if doc['development']['illustrative']: blockers.append('ILLUSTRATIVE')
    checks=[{'id':rid,'section':sec} for rid,sec,_ in all_records]
    return {'structure':'FAIL' if errors else 'PASS','errors':errors,'gaps':gaps,'blockers':sorted(set(blockers)),
            'readiness':'BLOCKED' if errors or blockers else 'READY_FOR_CODING','checks':checks,
            'runtime':doc['verification']['runtime']['status'],'uat':doc['verification']['uat']['status']}

def wrap(content,id_,slug,name,version):
    scope=content.get('bolumler',{}).get('1.1',{}).get('alanlar',{})
    edition=scope.get('sistem','UNKNOWN')
    if edition not in ('S/4HANA Cloud Public Edition','S/4HANA Cloud Private Edition'): edition='UNKNOWN'
    return {'schemaVersion':'2.0','development':{'id':id_,'slug':slug,'name':name,'edition':edition,'release':None,'illustrative':False},
      'handoff':{'version':version,'mode':'NEW','language':'tr','mockupCodeException':{'authorized':False,'receipt':None}},
      'content':content,'project':{'namingRules':[],'catalog':[]},
      'baseline':{'kind':'NEW','systemState':'NOT_CHECKED','reference':None,'payloadSha256':None,'verifiedAt':None,'evidenceRefs':[]},
      'approval':{'status':'DRAFT','payloadSha256':None,'by':None,'date':None,'receipt':None},
      'verification':{k:{'status':'UNKNOWN' if k in ('release','objects','authorization') else 'NOT_RUN','evidenceRefs':[]} for k in ('release','objects','authorization','runtime','uat')},
      'dependencies':[],'extensions':{},'assets':[],
      'mockup':{'url':None,'version':None,'status':'NOT_CREATED','evidenceRefs':[]},'delivery':None}

def flatten(value,path=''):
    result={}
    if isinstance(value,dict):
        if not value: result[path]={}
        for k,v in value.items(): result.update(flatten(v,path+'/'+str(k).replace('~','~0').replace('/','~1')))
    elif isinstance(value,list):
        if not value: result[path]=[]
        row_keys=[r.get('id') if isinstance(r,dict) else r[0] if isinstance(r,list) and r else None for r in value]
        stable=bool(value) and all(isinstance(k,str) and ID.fullmatch(k) for k in row_keys) and len(set(row_keys))==len(row_keys)
        if stable: result[path+'/@order']=row_keys
        for i,v in enumerate(value): result.update(flatten(v,path+'/'+(row_keys[i] if stable else str(i))))
    else: result[path]=value
    return result

def delta(before,after):
    if before['development']['id']!=after['development']['id']: raise Invalid('Separate development baselines required')
    left=flatten({'content':before['content'],'extensions':before['extensions'],'dependencies':before['dependencies']})
    right=flatten({'content':after['content'],'extensions':after['extensions'],'dependencies':after['dependencies']})
    rows=[]
    for path in sorted(set(left)|set(right)):
        bp,ap=path in left,path in right
        if bp and ap and left[path]==right[path]: continue
        rows.append({'path':path,'kind':'CHANGED' if bp and ap else 'ADDED' if ap else 'REMOVED',
                     'beforePresent':bp,'afterPresent':ap,'before':left.get(path),'after':right.get(path)})
    return {'status':'COMPUTED','baselinePayloadSha256':snapshot_sha(before),'targetPayloadSha256':snapshot_sha(after),'changes':rows}

def pointer(doc,path,value,remove=False):
    if not path.startswith('/') or path=='/': raise Invalid('Use a non-root JSON Pointer')
    parts=[p.replace('~1','/').replace('~0','~') for p in path[1:].split('/')]
    target=doc
    for key in parts[:-1]: target=target[int(key)] if isinstance(target,list) else target[key]
    key=parts[-1]
    if isinstance(target,list):
        if key=='-' and not remove: target.append(value)
        elif remove: del target[int(key)]
        else: target[int(key)]=value
    elif remove: del target[key]
    else: target[key]=value

def projection_rows(doc):
    # Every decoded leaf is preserved as path + JSON value. Display tabs are projections.
    return [(path,json.dumps(v,ensure_ascii=False,allow_nan=False)) for path,v in flatten(doc).items()]

def render_xlsx(doc,path,changes=None,validation=None):
    from openpyxl import Workbook
    from openpyxl.styles import Font,PatternFill,Alignment
    wb=Workbook(); wb.properties.creator=''; wb.properties.lastModifiedBy=''
    wb.properties.title=doc['development']['name']; wb.properties.description='FS-TS veri görünümü'
    summary=wb.active; summary.title='Özet'
    status=validation or validate(doc)
    for row in [('FS-TS',doc['development']['name']),('Geliştirme',doc['development']['id']),('Sürüm',doc['handoff']['version']),
                ('Durum',status['readiness']),('Kaynak','FS-TS.toon'),('Runtime',status['runtime']),('UAT',status['uat'])]: summary.append(row)
    for no,section in doc['content']['bolumler'].items():
        label=legacy.BOLUM.get(no,{}).get('kisa',no)
        ws=wb.create_sheet(re.sub(r'[\\/*?:\[\]]','-',no+' '+label)[:31])
        ws.append([legacy.BOLUM.get(no,{}).get('baslik',no)])
        if section.get('gecerli') is False: ws.append(['Uygulanmıyor',section['gerekce']]); continue
        field_labels={x['anahtar']:x['etiket'] for x in legacy.BOLUM.get(no,{}).get('alanlar',[])}
        for key,val in section.get('alanlar',{}).items(): ws.append([field_labels.get(key,key),val])
        if section.get('satirlar'):
            columns=legacy.BOLUM.get(no,{}).get('sutunlar',[])
            ws.append([c['ad'] for c in columns])
            for row in section['satirlar']: ws.append(row)
    if doc['extensions']:
        ws=wb.create_sheet('İlişkili kayıtlar'); ws.append(['Tür','ID','Kayıt'])
        for kind,rs in doc['extensions'].items():
            for row in rs: ws.append([kind,row['id'],json.dumps(row,ensure_ascii=False)])
    if changes is not None:
        ws=wb.create_sheet('Değişiklikler'); ws.append(['Durum',changes['status']])
        ws.append(['Yol','Tür','Önce mevcut','Sonra mevcut','Önce','Sonra'])
        for r in changes.get('changes',[]): ws.append([r[k] if isinstance(r[k],bool) else json.dumps(r[k],ensure_ascii=False) for k in ('path','kind','beforePresent','afterPresent','before','after')])
    ws=wb.create_sheet('Veri sözleşmesi'); ws.append(['Yol','JSON değeri'])
    for row in projection_rows(doc): ws.append(row)
    ws.sheet_state='hidden'
    for ws in wb:
        ws.freeze_panes='A2'; ws.sheet_view.showGridLines=False
        for c in ws[1]: c.font=Font(name='Calibri',size=11,bold=True,color='FFFFFF'); c.fill=PatternFill('solid',fgColor='18344D')
        for column in ws.columns:
            letter=column[0].column_letter
            ws.column_dimensions[letter].width=min(70,max(16,max(len(str(c.value or '')) for c in column)*0.7))
        for row in ws:
            for c in row:
                if isinstance(c.value,str): c.data_type='s'  # user text cannot become a formula
                c.alignment=Alignment(vertical='top',wrap_text=True)
            lines=max(1,max(sum(max(1,math.ceil(len(s)/max(12,ws.column_dimensions[c.column_letter].width))) for s in str(c.value or '').split('\n')) for c in row))
            ws.row_dimensions[row[0].row].height=min(409,16*lines+6)
    wb.save(path)
    # Remove library identity from OOXML application metadata in the exported copy.
    with zipfile.ZipFile(path) as z: entries={n:z.read(n) for n in z.namelist()}
    entries['docProps/app.xml']=b'<?xml version="1.0" encoding="UTF-8"?><Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/extended-properties"><Application>Spreadsheet</Application><AppVersion>1.0</AppVersion></Properties>'
    with zipfile.ZipFile(path,'w',zipfile.ZIP_DEFLATED) as z:
        for n,data in entries.items(): z.writestr(n,data)

def verify_xlsx(doc,path,validation=None):
    from openpyxl import load_workbook
    wb=load_workbook(path,data_only=False)
    actual=list(wb['Veri sözleşmesi'].values)[1:]
    if actual!=projection_rows(doc): raise Invalid('Excel source-data parity mismatch')
    if validation is not None and wb['Özet']['B4'].value!=validation['readiness']: raise Invalid('Excel readiness mismatch')
    for no,section in doc['content']['bolumler'].items():
        label=legacy.BOLUM.get(no,{}).get('kisa',no)
        ws=wb[re.sub(r'[\\/*?:\[\]]','-',no+' '+label)[:31]]
        visible=list(ws.values)
        if section.get('gecerli') is False:
            if visible[1][:2]!=('Uygulanmıyor',section['gerekce']): raise Invalid('Excel applicability mismatch')
            continue
        pos=1
        for _,value in section.get('alanlar',{}).items():
            if visible[pos][1]!=value: raise Invalid('Excel visible field mismatch: '+no)
            pos+=1
        if section.get('satirlar'):
            pos+=1
            for expected in section['satirlar']:
                if list(visible[pos][:len(expected)])!=expected: raise Invalid('Excel visible row mismatch: '+no)
                pos+=1
    for ws in wb:
        for row in ws:
            if any(c.data_type=='f' for c in row): raise Invalid('Unexpected workbook formula')
    return len(actual)

def safe_name(name):
    p=PurePosixPath(name)
    if not name or p.is_absolute() or '..' in p.parts or '\\' in name or ':' in name or str(p)!=name:
        raise Invalid('Unsafe archive/asset path')
    if any(part in ('.','') or part.upper().split('.')[0] in ('CON','PRN','AUX','NUL','COM1','LPT1') for part in p.parts):
        raise Invalid('Unsafe path component')
    return p

def asset_bytes(doc,asset,base):
    rel=safe_name(asset['path']); base=Path(base).resolve(); candidate=base.joinpath(*rel.parts)
    if candidate.is_symlink() or any(p.is_symlink() for p in candidate.parents if p!=base and base in p.parents): raise Invalid('Symlink asset refused')
    path=candidate.resolve()
    if not path.is_relative_to(base) or not path.is_file(): raise Invalid('Missing or escaped asset')
    if path.stat().st_size>32*1024*1024: raise Invalid('Asset exceeds 32 MiB')
    data=path.read_bytes()
    if digest(data)!=asset['sha256']: raise Invalid('Asset hash mismatch')
    suffix=path.suffix.lower()
    if asset['category']=='mockup':
        exception=doc['handoff']['mockupCodeException']
        if not exception['authorized'] or not exception['receipt']: raise Invalid('Executable mockup needs explicit receipt')
        if suffix!='.zip': raise Invalid('Mockup exception accepts a supplied mockup ZIP only')
        validate_mockup(data)
    elif suffix in ('.toon','.md','.csv','.xml'):
        text=data.decode('utf-8',errors='strict')
        if neutrality(text): raise Invalid('Forbidden source/producer in asset')
        if suffix=='.xml' and re.search(r'(?i)<!DOCTYPE|<script|\bon\w+\s*=',text): raise Invalid('Active XML content refused')
    elif suffix=='.png':
        # Preserve pixels, remove text/EXIF attribution metadata without touching input.
        if data[:8]!=b'\x89PNG\r\n\x1a\n': raise Invalid('Invalid PNG')
        result=data[:8]; pos=8
        while pos<len(data):
            if pos+12>len(data): raise Invalid('Truncated PNG')
            length=int.from_bytes(data[pos:pos+4],'big'); end=pos+12+length
            if end>len(data): raise Invalid('Truncated PNG chunk')
            tag=data[pos+4:pos+8]
            if tag not in (b'tEXt',b'iTXt',b'zTXt',b'eXIf'): result+=data[pos:end]
            pos=end
        data=result
    else: raise Invalid('Unsupported non-code asset format; use PNG or text contract')
    return asset['category']+'/'+str(rel),data

def validate_mockup(data):
    allowed={'.html','.js','.mjs','.css','.json','.xml','.png','.jpg','.jpeg','.webp','.svg','.txt','.md','.map'}
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            seen=set(); total=0
            for info in z.infolist():
                if info.is_dir(): continue
                safe_name(info.filename)
                if info.filename in seen or info.flag_bits&1 or (info.external_attr>>16)&0o170000==0o120000: raise Invalid('Unsafe mockup entry')
                seen.add(info.filename); total+=info.file_size
                if total>128*1024*1024 or Path(info.filename).suffix.lower() not in allowed: raise Invalid('Mockup contains unsupported code/binary')
                if re.search(r'(?i)(?:^|/)(?:node_modules|\.git|\.claude|\.codex|hooks)(?:/|$)|\.env(?:$|\.)|SKILL\.md',info.filename): raise Invalid('Mockup contains development/provider files')
                payload=z.read(info)
                if Path(info.filename).suffix.lower() in {'.html','.js','.mjs','.css','.json','.xml','.svg','.txt','.md','.map'} and PRODUCER.search(payload.decode('utf-8')): raise Invalid('Mockup producer attribution')
            if not any(n.endswith('.html') for n in seen): raise Invalid('Mockup has no HTML entry')
    except zipfile.BadZipFile as e: raise Invalid('Invalid mockup archive') from e

def package(doc,output,assets_root,baseline=None):
    import delivery
    return delivery.package(doc,output,assets_root)

def verify_zip(path,expected_name=None):
    import delivery
    return delivery.verify(path,expected_name)

def output(value): print(json.dumps(value,ensure_ascii=False,indent=None,allow_nan=False))

def main(argv=None):
    if hasattr(sys.stdout,'reconfigure'): sys.stdout.reconfigure(encoding='utf-8')
    p=argparse.ArgumentParser(description='Role-owned TOON specifications / handoff v3')
    sub=p.add_subparsers(dest='cmd',required=True)
    for name in ('init','migrate'):
        q=sub.add_parser(name); q.add_argument('--slug',required=True); q.add_argument('--version',required=True); q.add_argument('--output',required=True)
        if name=='migrate': q.add_argument('source')
        else:
            q.add_argument('--id',required=True); q.add_argument('--name',required=True); q.add_argument('--types',required=True); q.add_argument('--edition',required=True,choices=['S/4HANA Cloud Public Edition','S/4HANA Cloud Private Edition']); q.add_argument('--profile',default='standart',choices=['hafif','standart','tam'])
    q=sub.add_parser('guide'); q.add_argument('--section'); q.add_argument('--group'); q.add_argument('--types',default=''); q.add_argument('--profile',default='standart')
    q=sub.add_parser('inspect'); q.add_argument('source'); q.add_argument('--baseline')
    q=sub.add_parser('patch'); q.add_argument('source'); q.add_argument('--path',required=True); q.add_argument('--value'); q.add_argument('--remove',action='store_true')
    q=sub.add_parser('approve'); q.add_argument('source'); q.add_argument('--by',required=True); q.add_argument('--date',required=True); q.add_argument('--receipt',required=True); q.add_argument('--confirmed',action='store_true',required=True)
    q=sub.add_parser('delta'); q.add_argument('baseline'); q.add_argument('target'); q.add_argument('--output',required=True)
    q=sub.add_parser('context'); q.add_argument('source'); q.add_argument('--sections',required=True); q.add_argument('--output',required=True)
    q=sub.add_parser('handoff'); q.add_argument('source'); q.add_argument('--output',required=True); q.add_argument('--assets-root'); q.add_argument('--baseline')
    q=sub.add_parser('verify'); q.add_argument('source')
    for op in ('release-init','release-inspect','release-approve','confirm-reviews'):
        q=sub.add_parser(op);q.add_argument('source')
        if op in ('release-approve','confirm-reviews'):
            q.add_argument('--confirmed',action='store_true',required=True);q.add_argument('--receipt',required=True)
        if op=='release-approve': q.add_argument('--defaults',action='store_true')
    q=sub.add_parser('profile');q.add_argument('--collection',required=True)
    q=sub.add_parser('eval-request');q.add_argument('source');q.add_argument('--round',type=int,required=True);q.add_argument('--output',required=True);q.add_argument('--assets-root')
    q=sub.add_parser('eval-record');q.add_argument('source');q.add_argument('record')
    q=sub.add_parser('read-reference');q.add_argument('source');q.add_argument('--reference-id',required=True,action='append');q.add_argument('--allow-legacy',action='store_true')
    q=sub.add_parser('feedback');q.add_argument('source');q.add_argument('--category',required=True);q.add_argument('--defect',required=True)

    for op in ('release-upgrade','questions','check-plan','check-record','preflight','status'):
        q=sub.add_parser(op);q.add_argument('source')
        if op in ('check-plan','check-record','preflight','status'): q.add_argument('--assets-root')
        if op=='check-plan': q.add_argument('--final',action='store_true')
        if op=='check-record':
            q.add_argument('--check',required=True);q.add_argument('--input-sha256',required=True);q.add_argument('--status',required=True,choices=['PASS','FAIL'])
    q=sub.add_parser('handoff-batch');q.add_argument('manifest');q.add_argument('--output',required=True);q.add_argument('--batch-id',required=True)
    import profile_runtime
    profile_runtime.add_parsers(sub)

    a=p.parse_args(argv)
    if a.cmd=='read-reference':
        import source_pool
        reader=source_pool.Reader(load(a.source),allow_legacy=a.allow_legacy)
        if len(a.reference_id)==1:
            identity=a.reference_id[0]
            output({'reference_id':identity,'legacy_read_only':reader.legacy_read_only,'value':reader.resolve(identity)})
        else:
            output({'values':reader.resolve_many(a.reference_id),'legacy_read_only':reader.legacy_read_only})
        return
    if a.cmd=='work-profiles':
        import work_profiles
        output(work_profiles.catalog());return
    if a.cmd=='profile':
        import delivery
        if a.collection not in delivery.SCHEMA['properties']: raise Invalid('Unknown collection')
        output(delivery.SCHEMA['properties'][a.collection]);return
    if a.cmd=='guide':
        if not a.section and not a.group: raise Invalid('Select section or group; no full catalog dump')
        selected=[b for b in legacy.T['bolumler'] if (not a.section or b['no']==a.section) and (not a.group or b['grup']==a.group) and legacy.bolum_durumu(b,a.types.split(',') if a.types else [],a.profile)=='gecerli']
        output({'sections':[{'id':b['no'],'title':b['baslik'],'purpose':b['amac'],'fields':b['alanlar'],'columns':b['sutunlar'],'minRows':b['en_az_satir']} for b in selected]}); return
    if a.cmd=='migrate':
        original=Path(a.source).read_bytes()
        content=strict_json(original.decode('utf-8'))
        ensure_numbers(content)
        scope=content.get('bolumler',{}).get('1.1',{}).get('alanlar',{})
        if not scope.get('gel_id') or not scope.get('baslik'): raise Invalid('Legacy development identity missing')
        doc=wrap(content,scope['gel_id'],a.slug,scope['baslik'],a.version)
        if schema_check(doc): raise Invalid('Legacy document shape invalid')
        save(doc,a.output)
        if Path(a.source).read_bytes()!=original: raise Invalid('Original modified')
        output({'path':a.output,'originalSha256':digest(original),'contentPreserved':True,'readiness':'BLOCKED'}); return
    if a.cmd=='init':
        types=a.types.split(',')
        if any(t not in legacy.TUR_AD for t in types): raise Invalid('Unknown RICEF type')
        sections={}
        for b in legacy.T['bolumler']:
            if b.get('turetilmis') or legacy.bolum_durumu(b,types,a.profile)!='gecerli': continue
            body={}
            if b['alanlar']: body['alanlar']={f['anahtar']:'BİLGİ BEKLİYOR (OPEN-01)' for f in b['alanlar']}
            if b['sutunlar']: body['satirlar']=[]
            sections[b['no']]=body
        sections['1.1']['alanlar'].update(gel_id=a.id,baslik=a.name,sistem=a.edition)
        sections['7.3']={'satirlar':[['OPEN-01','Eksik girdi ve kararlar','1.1, 2.2, 3.5','Süreç sahibi','—','Açık','—','Karar bekleyen konu']]}
        content={'sema':'1.0','meta':{'surum':a.version,'tarih':'—','hazirlayan':'—','durum':'Taslak','musteri':'—'},'turler':types,'profil':a.profile,'bolumler':sections,'surum_gecmisi':[],'onaylar':[]}
        doc=wrap(content,a.id,a.slug,a.name,a.version); save(doc,a.output); output({'path':a.output,'readiness':'BLOCKED'}); return
    if a.cmd=='verify': output(verify_zip(a.source)); return
    if a.cmd=='delta': save(delta(load(a.baseline),load(a.target)),a.output); output({'path':a.output}); return
    if a.cmd=='handoff-batch':
        import handoff3 as h
        manifest=load(a.manifest);base=Path(a.manifest).resolve().parent;items=[];paths=[]
        if set(manifest)!={'developments'}: raise Invalid('Use a batch manifest with developments only')
        for entry in manifest['developments']:
            if set(entry)!={'workspace','assets_root'}: raise Invalid('Batch entry needs workspace and assets_root')
            def scoped(raw):
                result=base.joinpath(*safe_name(raw).parts)
                if result.is_symlink() or not result.resolve().is_relative_to(base): raise Invalid('Batch path escaped its input root')
                return result.resolve()
            source=scoped(entry['workspace']);assets=scoped(entry['assets_root'])
            items.append({'doc':load(source),'assets_root':assets});paths.append(source)
        if len(set(paths))!=len(paths): raise Invalid('Batch repeats a workspace')
        from contextlib import ExitStack
        import workspace_lock
        with ExitStack() as locks:
            for source in sorted(paths,key=str):locks.enter_context(workspace_lock.held(source))
            for item,source in zip(items,paths):item['doc']=load(source)
            result=h.batch(items,a.output,a.batch_id)
            for item,source in zip(items,paths):save(item['doc'],source,True)
            import profile_runtime
            import dispatch_registry
            for item,source in zip(items,paths):
                if profile_runtime.managed(item['doc']):dispatch_registry.finish_selection(item['doc'],source)
        output(result);return

    import workspace_lock
    with workspace_lock.held(a.source):
        return _workspace_command(a)


def _workspace_command(a):
    doc=load(a.source)
    if a.cmd.startswith('work-'):
        import profile_runtime
        result,changed=profile_runtime.run(a,doc)
        if changed:
            save(doc,a.source,True)
            profile_runtime.after_save(a,doc,result)
        output(result);return
    if a.cmd in ('release-upgrade','questions','check-plan','check-record','preflight','status'):
        import handoff3 as h
        if a.cmd=='questions': output({'questions':h.consultant_questions(doc),'technical_questions_sent_to_consultant':False});return
        if a.cmd in ('preflight','status'):
            import preflight
            output(preflight.inspect(doc,a.assets_root or Path(a.source).parent) if a.cmd=='preflight' else preflight.status(doc,a.assets_root or Path(a.source).parent));return
        if a.cmd=='check-plan': output(h.check_plan(doc,a.final,a.assets_root or Path(a.source).parent));return
        if a.cmd=='release-upgrade': result=h.upgrade(doc)
        else:
            h.record_check(doc,a.check,a.input_sha256,a.status,a.assets_root or Path(a.source).parent);result={'check':a.check,'status':a.status}
        save(doc,a.source,True);output(result);return
    if a.cmd in ('release-init','release-inspect','release-approve','confirm-reviews','eval-request','eval-record','feedback'):
        import delivery
        if a.cmd=='release-init':
            delivery.release_init(doc)
            doc['delivery']['control']['work_profile_required']=True
        elif a.cmd=='release-inspect': output(delivery.evaluate(doc));return
        elif a.cmd=='release-approve':
            spec,control=delivery.get_state(doc)
            if delivery.shape_and_profile(spec): raise Invalid('Close deterministic defects before snapshot approval')
            if a.defaults: control.update(approved_defaults_sha256=delivery.defaults_sha(spec),defaults_approval_receipt=a.receipt)
            else: control.update(approved_spec_sha256=delivery.revision(spec),approval_receipt=a.receipt)
        elif a.cmd=='confirm-reviews':
            import source_pool
            _,control=delivery.get_state(doc);control.update(review_execution_confirmed=True,review_confirmation_receipt=a.receipt,review_confirmation_protocol=source_pool.PROTOCOL)
        elif a.cmd=='eval-request':
            import source_pool
            packets=delivery.eval_request(doc,a.round,a.assets_root or Path(a.source).parent)
            source_pool.save_packets(packets,a.output,doc['delivery']['control'])
        elif a.cmd=='eval-record': delivery.record_round(doc,load(a.record))
        else: delivery.feedback(doc,a.defect,a.category)
        save(doc,a.source,True);output({'path':a.source,'operation':a.cmd,'deliveryGate':delivery.evaluate(doc)['decision']});return
    if a.cmd=='handoff':
        import delivery
        result=delivery.package(doc,a.output,a.assets_root or Path(a.source).parent)
        save(doc,a.source,True)
        import profile_runtime
        if profile_runtime.managed(doc):
            import dispatch_registry
            dispatch_registry.finish_selection(doc,a.source)
        output(result);return
    if a.cmd=='patch':
        prior_control=doc.get('delivery',{}).get('control',{}) if isinstance(doc.get('delivery'),dict) else {}
        protected={key:copy.deepcopy(prior_control[key]) for key in ('work_profile','work_profile_required') if key in prior_control}
        if a.path.startswith('/approval'): raise Invalid('Use explicit approve command')
        if a.path.startswith('/delivery/control/work_profile'):
            raise Invalid('Use work-select/begin/finish; the episode ledger cannot be reset by patch')
        if doc.get('delivery') is not None:
            import work_profiles
            work_profiles.admit(doc,'write')
        if doc.get('delivery') is not None and a.path.startswith('/content'): raise Invalid('Imported legacy content is reference-only after release-init; edit delivery.spec')
        pointer(doc,a.path,json.loads(a.value) if a.value is not None else None,a.remove)
        after_control=doc.get('delivery',{}).get('control',{}) if isinstance(doc.get('delivery'),dict) else {}
        if protected!={key:after_control[key] for key in ('work_profile','work_profile_required') if key in after_control}:
            raise Invalid('Patch cannot replace or remove a managed episode or selection requirement')
        if doc.get('delivery') is not None and a.path.startswith('/delivery/spec'):
            import delivery
            delivery.invalidate(doc,defaults=a.path.startswith('/delivery/spec/functional_defaults'))
        if schema_check(doc): raise Invalid('Patch violates data contract')
        doc['approval']={'status':'DRAFT','payloadSha256':None,'by':None,'date':None,'receipt':None}
        save(doc,a.source,True); output({'path':a.source,'approval':'DRAFT'}); return
    if a.cmd=='approve':
        if not re.fullmatch(r'\d{4}-\d{2}-\d{2}',a.date): raise Invalid('Use ISO date supplied with approval')
        date.fromisoformat(a.date)
        doc['approval']={'status':'APPROVED_SNAPSHOT','payloadSha256':snapshot_sha(doc),'by':a.by,'date':a.date,'receipt':a.receipt}
        save(doc,a.source,True); output({'path':a.source,'snapshotSha256':snapshot_sha(doc),'recordedNotAuthenticated':True}); return
    if a.cmd=='context':
        if doc.get('delivery') is not None:
            import delivery
            spec,_=delivery.get_state(doc)
            if delivery.shape_and_profile(spec): raise Invalid('Complete public profile before selecting context')
            packet=delivery.context_packet(spec,a.sections.split(','));save(packet,a.output)
            output({'path':a.output,'includedRecords':len(packet['included_ids']),'privateControlIncluded':False});return
        if schema_check(doc): raise Invalid('Invalid input')
        selected=set(a.sections.split(','))|{'1.1','7.3'}
        idmap={rid:sec for rid,sec,_ in records(doc)}
        changed=True
        while changed:
            changed=False
            candidates=[(sec,body) for sec,body in doc['content']['bolumler'].items()]
            candidates += [('extensions/'+sec,rows) for sec,rows in doc['extensions'].items()]
            for sec,body in candidates:
                if sec in selected:
                    if sec.startswith('extensions/'):
                        refs=[ref for row in body for ref in row.get('refs',[])]
                    else:
                        text=json.dumps(body,ensure_ascii=False)
                        refs=[m.group(0) for m in legacy.KIMLIK_RE.finditer(text)]
                    for ref in refs:
                        if ref not in idmap: raise Invalid('Unresolved context ref: '+ref)
                        if idmap[ref] not in selected: selected.add(idmap[ref]); changed=True
        result=copy.deepcopy(doc)
        result['content']['bolumler']={k:v for k,v in doc['content']['bolumler'].items() if k in selected}
        result['extensions']={k:v for k,v in doc['extensions'].items() if 'extensions/'+k in selected}
        packet={'kind':'SELECTED_CONTEXT','authoritative':False,'sourcePayloadSha256':snapshot_sha(doc),'includedSections':sorted(selected),'omittedSections':sorted(set(doc['content']['bolumler'])-selected),'data':result}
        save(packet,a.output); output({'path':a.output,'sections':sorted(selected)}); return
    baseline=load(a.baseline) if a.baseline else None
    if a.cmd=='inspect':
        result=validate(doc,baseline); result['checks']={'recordCount':len(result['checks'])}; result['errors']=result['errors'][:20]; result['gaps']=result['gaps'][:20]; output(result); return
    if a.cmd=='handoff': output(package(doc,a.output,a.assets_root or Path(a.source).parent,baseline))

if __name__=='__main__':
    try: main()
    except (Invalid,OSError,KeyError,ValueError,TypeError,ImportError,subprocess.SubprocessError) as e:
        print(json.dumps({'status':'BLOCKED' if str(e).startswith('BLOCKED:') else 'FAIL','message':str(e)},ensure_ascii=False)); sys.exit(1)
