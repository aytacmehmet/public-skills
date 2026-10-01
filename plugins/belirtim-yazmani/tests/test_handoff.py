# SPDX-License-Identifier: GPL-3.0-only
import copy,io,json,os,subprocess,sys,tempfile,unittest,zipfile
from pathlib import Path
ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'scripts'))
import bv2 as bv

class HandoffTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls): cls.base=bv.load(ROOT/'assets/example.toon')
    def setUp(self):
        self.doc=copy.deepcopy(self.base); self.tmp=tempfile.TemporaryDirectory(); self.dir=Path(self.tmp.name)
    def tearDown(self): self.tmp.cleanup()
    def codes(self,doc=None,base=None): return {e['code'] for e in bv.validate(doc or self.doc,base)['errors']}
    def write(self,doc=None,name='source.toon'):
        p=self.dir/name; bv.save(doc or self.doc,p); return p
    def cli(self,*args):
        r=subprocess.run([sys.executable,str(ROOT/'scripts/bv2.py'),*map(str,args)],cwd=self.dir,text=True,encoding='utf-8',capture_output=True)
        return r,json.loads(r.stdout)
    def signed(self,doc):
        doc['approval']={'status':'APPROVED_SNAPSHOT','payloadSha256':bv.snapshot_sha(doc),'by':'Fixture-only actor','date':'2026-09-30','receipt':'TEST_ONLY_NOT_A_REAL_APPROVAL'}
        return doc
    def mutate_zip(self,path,mutation):
        with zipfile.ZipFile(path) as z: files={n:z.read(n) for n in z.namelist()}
        mutation(files)
        with zipfile.ZipFile(path,'w',zipfile.ZIP_DEFLATED) as z:
            for n,b in files.items(): z.writestr(n,b)
    def test_official_codec_types_unicode_quotes(self):
        value={'records':[{'id':'01','value':'İstanbul, "ABC"','number':'1.250'},{'id':'02','value':'satır\niki','number':'0.001'}],'nil':None,'ok':True,'empty':[],'object':{},'nested':[{'x':[1,2]}]}
        self.assertEqual(json.loads(bv.codec('decode',bv.encode(value))),value)
    def test_invalid_array_count(self):
        with self.assertRaises(bv.Invalid): bv.codec('decode','rows[2]{id,x}:\n  a,b\n')
    def test_invalid_row_width(self):
        with self.assertRaises(bv.Invalid): bv.codec('decode','rows[1]{id,x}:\n  a,b,c\n')
    def test_duplicate_toon_key(self):
        with self.assertRaises(bv.Invalid): bv.codec('decode','a: 1\na: 2\n')
    def test_indentation_rejected(self):
        with self.assertRaises(bv.Invalid): bv.codec('decode','a:\n\tb: 2\n')
    def test_large_integer_requires_string(self):
        with self.assertRaises(bv.Invalid): bv.encode({'id':9007199254740993})
        self.assertIn('9007199254740993',bv.encode({'id':'9007199254740993'}))
    def test_nonfinite_rejected(self):
        with self.assertRaises(bv.Invalid): bv.encode({'x':float('nan')})
    def test_json_duplicates_and_precision(self):
        with self.assertRaises(bv.Invalid): bv.strict_json('{"a":1,"a":2}')
        with self.assertRaises(bv.Invalid): bv.strict_json('{"a":0.1234567890123456789012345}')
    def test_valid_draft_is_blocked_and_reference_rows_are_not_definitions(self):
        result=bv.validate(self.doc)
        self.assertEqual(result['structure'],'PASS'); self.assertEqual(result['readiness'],'BLOCKED')
        self.assertIn('ILLUSTRATIVE',result['blockers']); self.assertIn('SAP_OBJECT_PROVENANCE',result['blockers'])
        self.assertNotIn('DUPLICATE_ID',self.codes())
    def test_schema_rejects_wrong_type_and_extra_field(self):
        self.doc['handoff']['version']=3; self.doc['extra']=True
        self.assertIn('SCHEMA',self.codes())
    def test_duplicate_definition(self):
        self.doc['extensions']={'ui':[{'id':'REQ-01','refs':[]}]}
        self.assertIn('DUPLICATE_ID',self.codes())
    def test_unresolved_reference(self):
        self.doc['extensions']={'ui':[{'id':'UIX-01','refs':['NOPE-01']}]}
        self.assertIn('UNRESOLVED_REF',self.codes())
    def test_foreign_mutation_requires_separate_handoff(self):
        self.doc['extensions']={'ui':[{'id':'UIX-01','refs':[],'developmentId':'RMP'}]}
        self.assertIn('FOREIGN_MUTATION',self.codes())
    def test_reference_dependency_allowed_mutation_rejected(self):
        dep={'developmentId':'RMP','version':'1.0.0','contract':'Contract reference','operation':'REFERENCE','status':'UNKNOWN','evidenceRefs':[]}
        self.doc['dependencies']=[dep]
        self.assertEqual(self.codes(),set()); self.assertIn('DEPENDENCY:RMP',bv.validate(self.doc)['blockers'])
        dep['operation']='CHANGE'; self.assertIn('SCHEMA',self.codes())
    def test_identity_mismatch(self):
        self.doc['development']['id']='OTHER'; self.assertIn('SINGLE_DEVELOPMENT',self.codes())
    def test_no_access_is_not_absence(self):
        self.doc['handoff']['mode']='UPDATE'; self.doc['baseline'].update(kind='APPROVED_HANDOFF',systemState='UNKNOWN')
        self.assertIn('ABSENCE_NOT_PROVEN',self.codes())
    def test_invalid_system_baseline_presence(self):
        self.doc['handoff']['mode']='UPDATE'; self.doc['baseline'].update(kind='SYSTEM',systemState='ABSENT_VERIFIED')
        self.assertIn('SYSTEM_BASELINE',self.codes())
    def test_stale_approval_rejected(self):
        self.signed(self.doc); self.doc['development']['name']='Changed after approval'
        self.assertIn('APPROVAL_SNAPSHOT',self.codes())
    def test_delta_stable_ids_null_and_empty_types(self):
        before=copy.deepcopy(self.doc); after=copy.deepcopy(self.doc)
        after['content']['bolumler']['2.2']['satirlar'].reverse()
        changes=bv.delta(before,after)['changes']
        self.assertEqual([r['path'] for r in changes],['/content/bolumler/2.2/satirlar/@order'])
        before['extensions']={'test':[{'id':'EXT-01','refs':[],'optional':None,'data':{}}]}
        after['extensions']={'test':[{'id':'EXT-01','refs':[],'data':'{}'}]}
        rows=bv.delta(before,after)['changes']; removed=next(r for r in rows if r['path'].endswith('/optional'))
        self.assertTrue(removed['beforePresent']); self.assertFalse(removed['afterPresent'])
        self.assertIsNone(removed['before']); self.assertTrue(any(r['before']=={} and r['after']=='{}' for r in rows))
    def test_delta_foreign_identity_rejected(self):
        after=copy.deepcopy(self.doc); after['development']['id']='OTHER'
        with self.assertRaises(bv.Invalid): bv.delta(self.doc,after)
    def test_producer_field_model_setting_and_inline_implementation(self):
        for row,code in [({'id':'EXT-01','refs':[],'llmModel':'any'},'PRODUCER_FIELD'),({'id':'EXT-01','refs':[],'note':'Generated by Claude'},'PRODUCER_ATTRIBUTION'),({'id':'EXT-01','refs':[],'note':'METHOD save.\nENDMETHOD.'},'IMPLEMENTATION')]:
            self.doc['extensions']={'test':[row]}; self.assertIn(code,self.codes())
    def test_material_model_field_is_not_llm_metadata(self):
        self.doc['extensions']={'domain':[{'id':'MAT-01','refs':[],'model':'mill-type-a'}]}
        self.assertEqual(self.codes(),set())
    def test_excel_visible_and_source_mutation_rejected(self):
        from openpyxl import load_workbook
        p=self.dir/'book.xlsx'; bv.render_xlsx(self.doc,p); self.assertGreater(bv.verify_xlsx(self.doc,p),0)
        wb=load_workbook(p); ws=next(w for w in wb if w.title.startswith('2.2 ')); ws['B2']='WRONG'; wb.save(p)
        with self.assertRaises(bv.Invalid): bv.verify_xlsx(self.doc,p)
    def test_formula_payload_stays_text(self):
        from openpyxl import load_workbook
        self.doc['content']['bolumler']['1.1']['alanlar']['baslik']='=1+1'
        p=self.dir/'book.xlsx'; bv.render_xlsx(self.doc,p)
        wb=load_workbook(p); self.assertTrue(all(c.data_type!='f' for w in wb for row in w for c in row))
    def test_code_asset_and_traversal_rejected(self):
        for path in ['../secret.md','/absolute.md','a\\b.md','C:/x.md']:
            with self.assertRaises(bv.Invalid): bv.safe_name(path)
        p=self.dir/'program.abap'; p.write_text('REPORT zfoo.')
        asset={'path':'program.abap','category':'contracts','description':'Source','sha256':bv.digest(p.read_bytes())}
        with self.assertRaises(bv.Invalid): bv.asset_bytes(self.doc,asset,self.dir)
    def test_missing_packaged_evidence_rejected(self):
        self.doc['verification']['release']={'status':'VERIFIED','evidenceRefs':['evidence/missing.md']}
        with self.assertRaises(bv.Invalid): bv.package(self.doc,self.dir,self.dir)
    def test_patch_invalidates_approval(self):
        self.signed(self.doc); source=self.write()
        r,result=self.cli('patch',source,'--path','/content/meta/musteri','--value','"New Customer"')
        self.assertEqual(r.returncode,0); self.assertEqual(bv.load(source)['approval']['status'],'DRAFT')
    def test_migration_content_original_and_unknowns(self):
        source=self.dir/'old.json'; original=json.dumps(self.doc['content'],ensure_ascii=False).encode();source.write_bytes(original)
        target=self.dir/'migrated.toon'
        r,result=self.cli('migrate',source,'--slug','demo-satin-alma','--version','2.0.0','--output',target)
        self.assertEqual(r.returncode,0); self.assertEqual(source.read_bytes(),original)
        new=bv.load(target);self.assertEqual(new['content'],self.doc['content']);self.assertEqual(new['approval']['status'],'DRAFT');self.assertIsNone(new['development']['release'])
    def test_context_reference_closure_and_not_authoritative(self):
        source=self.write(); target=self.dir/'context.toon'
        r,result=self.cli('context',source,'--sections','3.7','--output',target)
        self.assertEqual(r.returncode,0);packet=bv.load(target)
        self.assertFalse(packet['authoritative']);self.assertEqual(packet['kind'],'SELECTED_CONTEXT')
        self.assertIn('3.3',packet['includedSections']);self.assertIn('7.3',packet['includedSections'])
    def test_guide_works_from_foreign_cwd(self):
        r,result=self.cli('guide','--section','3.4','--types','U','--profile','standart')
        self.assertEqual(r.returncode,0);self.assertEqual([s['id'] for s in result['sections']],['3.4'])
    def test_init_never_invents_business_decisions(self):
        target=self.dir/'new.toon'
        r,result=self.cli('init','--id','GEL-DEMO-01','--name','Fixture','--slug','fixture','--version','1.0.0','--types','U','--edition','S/4HANA Cloud Public Edition','--output',target)
        self.assertEqual(r.returncode,0);d=bv.load(target)
        self.assertIn('BEKLİYOR',d['content']['bolumler']['1.1']['alanlar']['yaklasim']);self.assertEqual(d['approval']['status'],'DRAFT')

if __name__=='__main__': unittest.main(verbosity=2)
